#!/usr/bin/env python3
"""pdfstruct.export - turn an existing <name>.raw.json into another output format.

Never extracts or OCRs by itself. For every document:
  1. <output>/<name>.raw.json exists  -> it is used as is (the PDF is not opened)
  2. otherwise                        -> the normal extraction pipeline runs once
  3. the target format is generated from the raw JSON

These are generic document exports (pages and line blocks). Semantic fields
(vocabulary entries, invoice fields, ...) still belong in a project parser.

Usage:
    python -m pdfstruct.export --format html --input <pdf file or folder> --output <folder>
"""
from __future__ import annotations

import argparse
import contextlib
import csv
import html
import io
import json
import os
import re
import sqlite3
import sys
from pathlib import Path

from . import extract

FORMATS = {"html": ".html", "txt": ".txt", "md": ".md", "csv": ".csv", "xlsx": ".xlsx",
           "docx": ".docx", "jsonl": ".jsonl", "sqlite": ".sqlite"}
RAW_SUFFIX = ".raw.json"
BLOCK_COLUMNS = ["source_file", "page", "block_index", "text",
                 "x1", "y1", "x2", "y2", "method", "confidence"]
PAGE_COLUMNS = ["source_file", "page", "method", "width", "height", "rotation",
                "text_layer_score", "ocr_confidence", "needs_review", "warnings",
                "chars", "blocks", "text"]
XLSX_CELL_LIMIT = 32767

# pdfium marks a hyphen at a line break with one of these.
_HYPHEN_MARKERS = re.compile("[￾\x02]")
# Not allowed in XML (docx/xlsx) and useless in readable text.
_CONTROL_CHARS = re.compile("[\x00-\x08\x0b-\x1f\x7f-\x9f\ud800-\udfff￿]")
_MD_LINE_START = re.compile(r"^(\s*)(#{1,6}\s|>|[-+*]\s|\d+[.)]\s|={3,}\s*$|-{3,}\s*$|```|~~~|\|)")


def clean(text: str | None) -> str:
    """Text for readable outputs: hyphen markers become '-', control characters go."""
    return _CONTROL_CHARS.sub("", _HYPHEN_MARKERS.sub("-", text or "").replace("\r", ""))


def source_name(raw: dict) -> str:
    return (raw.get("source") or {}).get("file_name") or Path(raw.get("source_file") or "").name


def page_score(page: dict):
    return (page.get("text_layer") or {}).get("score")


def iter_blocks(raw: dict):
    """One flat record per line block, in page and extraction order."""
    name = source_name(raw)
    for page in raw["pages"]:
        for index, block in enumerate(page.get("blocks") or []):
            bbox = list(block.get("bbox") or [])[:4] + [None] * 4
            yield {"source_file": name, "page": page["page"],
                   "block_index": block.get("order", index), "text": block.get("text", ""),
                   "x1": bbox[0], "y1": bbox[1], "x2": bbox[2], "y2": bbox[3],
                   "method": page.get("method"), "confidence": block.get("confidence")}


def atomic_text(path: Path, text: str, encoding: str = "utf-8") -> None:
    tmp = path.with_name(path.name + ".tmp")
    with tmp.open("w", encoding=encoding, newline="") as out:
        out.write(text)
    os.replace(tmp, path)


def fmt_num(value, digits: int = 2) -> str:
    return "n/a" if value is None else format(value, f".{digits}f")


# ---------------------------------------------------------------- txt / md

def export_txt(raw: dict, path: Path) -> None:
    total = raw["page_count"]
    parts = []
    for page in raw["pages"]:
        parts.append(f"===== Page {page['page']} / {total} =====\n\n"
                     f"{clean(page.get('text')).strip()}\n")
    atomic_text(path, "\n".join(parts))


def export_md(raw: dict, path: Path) -> None:
    out = [f"# {source_name(raw)}", ""]
    for page in raw["pages"]:
        out += [f"## Page {page['page']}", ""]
        lines = clean(page.get("text")).strip().split("\n")
        for number, line in enumerate(lines):
            line = _MD_LINE_START.sub(lambda m: m.group(1) + "\\" + m.group(2), line.rstrip())
            # two trailing spaces = hard line break, so the line order survives rendering
            last = number == len(lines) - 1 or not line or not lines[number + 1].strip()
            out.append(line if last else line + "  ")
        out.append("")
    atomic_text(path, "\n".join(out))


# ---------------------------------------------------------------- html

_HTML_CSS = """
:root{--bg:#f6f6f4;--card:#fff;--ink:#1c1c1a;--muted:#6b6b66;--line:#e2e2dd;--accent:#1f5fbf;
--warn-bg:#fff3cd;--warn-ink:#7a5200;--ocr-bg:#e6f0ff;--ocr-ink:#1f4a8f;--ok-bg:#e7f4ea;--ok-ink:#216b35}
@media(prefers-color-scheme:dark){:root{--bg:#161614;--card:#20201d;--ink:#ecece6;--muted:#a3a39b;
--line:#34342f;--accent:#8ab4ff;--warn-bg:#4a3a10;--warn-ink:#ffd978;--ocr-bg:#1d2f4d;--ocr-ink:#a9c8ff;
--ok-bg:#1c3a25;--ok-ink:#9bd9ab}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 "Segoe UI",system-ui,sans-serif}
main{max-width:980px;margin:0 auto;padding:16px}
h1{font-size:20px;margin:8px 0;overflow-wrap:anywhere}
h2{font-size:16px;margin:0}
a{color:var(--accent)}
.meta,.head{display:flex;flex-wrap:wrap;gap:6px;align-items:center}
.chip{font-size:12px;padding:1px 8px;border-radius:10px;border:1px solid var(--line);color:var(--muted);white-space:nowrap}
.chip.native{background:var(--ok-bg);color:var(--ok-ink);border-color:transparent}
.chip.ocr{background:var(--ocr-bg);color:var(--ocr-ink);border-color:transparent}
.chip.review,.chip.error{background:var(--warn-bg);color:var(--warn-ink);border-color:transparent}
nav{margin:12px 0;display:flex;flex-wrap:wrap;gap:4px}
nav a{font-size:12px;min-width:30px;text-align:center;padding:2px 6px;border:1px solid var(--line);
border-radius:4px;text-decoration:none;background:var(--card)}
nav a.review{background:var(--warn-bg);color:var(--warn-ink)}
section{background:var(--card);border:1px solid var(--line);border-radius:6px;margin:12px 0;padding:12px 14px}
.head{padding-bottom:8px;margin-bottom:8px;border-bottom:1px solid var(--line)}
.warn{font-size:13px;color:var(--warn-ink);background:var(--warn-bg);padding:4px 8px;border-radius:4px;margin-bottom:8px}
.b{white-space:pre-wrap;overflow-wrap:anywhere;padding:1px 4px;border-radius:3px}
.b:hover{background:var(--bg)}
.b.low{background:var(--warn-bg)}
.b small{color:var(--muted);font-size:11px;margin-left:8px}
.empty{color:var(--muted);font-style:italic}
"""


def export_html(raw: dict, path: Path) -> None:
    esc = html.escape
    name = source_name(raw)
    review = set(raw.get("review_pages") or [])
    low_conf = extract.DEFAULT_CONFIG["quality"]["ocr_review_confidence"]
    out = ["<!doctype html>", '<html lang="und"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width,initial-scale=1">',
           f"<title>{esc(name)}</title>", f"<style>{_HTML_CSS}</style></head><body><main>",
           f"<h1>{esc(name)}</h1>", '<div class="meta">',
           f'<span class="chip">pages: {raw["page_count"]}</span>',
           f'<span class="chip">method: {esc(str(raw.get("extraction_method")))}</span>',
           f'<span class="chip">text layer score: {fmt_num(raw.get("text_layer_score"))}</span>',
           f'<span class="chip">OCR used: {"yes" if raw.get("ocr_used") else "no"}</span>',
           f'<span class="chip{" review" if review else ""}">review pages: {len(review)}</span>',
           "</div>", "<nav>"]
    for page in raw["pages"]:
        number = page["page"]
        cls = ' class="review"' if number in review else ""
        out.append(f'<a{cls} href="#p{number}">{number}</a>')
    out.append("</nav>")
    for page in raw["pages"]:
        number, method = page["page"], str(page.get("method"))
        blocks = page.get("blocks") or []
        out.append(f'<section id="p{number}"><div class="head"><h2>Page {number}</h2>'
                   f'<span class="chip {esc(method)}">method: {esc(method)}</span>'
                   f'<span class="chip">score: {fmt_num(page_score(page))}</span>')
        if page.get("ocr_confidence") is not None:
            out.append(f'<span class="chip">confidence: {fmt_num(page["ocr_confidence"])}</span>')
        out.append(f'<span class="chip">blocks: {len(blocks)}</span>')
        if page.get("needs_review"):
            out.append('<span class="chip review">needs review</span>')
        out.append("</div>")
        for warning in page.get("warnings") or []:
            out.append(f'<div class="warn">{esc(clean(str(warning)))}</div>')
        if blocks:
            for index, block in enumerate(blocks):
                conf = block.get("confidence")
                bbox = ", ".join(str(v) for v in block.get("bbox") or [])
                title = f"block {block.get('order', index)} | bbox {bbox}"
                small = ""
                if conf is not None:
                    title += f" | confidence {conf:.2f}"
                    small = f"<small>{conf:.2f}</small>"
                cls = "b low" if conf is not None and conf < low_conf else "b"
                out.append(f'<div class="{cls}" title="{esc(title)}">'
                           f'{esc(clean(block.get("text")))}{small}</div>')
        elif clean(page.get("text")).strip():
            out.append(f'<div class="b">{esc(clean(page.get("text")).strip())}</div>')
        else:
            out.append('<div class="empty">(no text on this page)</div>')
        out.append("</section>")
    out.append("</main></body></html>")
    atomic_text(path, "\n".join(out) + "\n")


# ---------------------------------------------------------------- csv / jsonl

def export_csv(raw: dict, path: Path) -> None:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=BLOCK_COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    for record in iter_blocks(raw):
        writer.writerow({k: "" if v is None else v for k, v in record.items()})
    # BOM so that Excel opens umlauts correctly; read it back with encoding="utf-8-sig"
    atomic_text(path, buffer.getvalue(), encoding="utf-8-sig")


def export_jsonl(raw: dict, path: Path) -> None:
    lines = []
    for record in iter_blocks(raw):
        bbox = [record.pop(key) for key in ("x1", "y1", "x2", "y2")]
        record["bbox"] = bbox
        lines.append(json.dumps(record, ensure_ascii=False))
    atomic_text(path, "".join(line + "\n" for line in lines))


# ---------------------------------------------------------------- xlsx

def _xlsx_sheet(workbook, title: str, columns: list[str], widths: dict, rows) -> None:
    from openpyxl.styles import Alignment, Font

    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    text_col = columns.index("text") + 1
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            # extracted text is never a formula, even if it starts with "="
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        sheet.cell(row=sheet.max_row, column=text_col).alignment = wrap
    for index, column in enumerate(columns, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = widths.get(column, 12)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_xlsx(raw: dict, path: Path) -> None:
    from openpyxl import Workbook

    name = source_name(raw)
    workbook = Workbook()
    workbook.remove(workbook.active)

    def cell_text(value) -> str:
        return clean(value)[:XLSX_CELL_LIMIT]

    pages = ([name, p["page"], p.get("method"), p.get("width"), p.get("height"),
              p.get("rotation"), page_score(p), p.get("ocr_confidence"),
              bool(p.get("needs_review")), cell_text("; ".join(map(str, p.get("warnings") or []))),
              len(clean(p.get("text"))), len(p.get("blocks") or []), cell_text(p.get("text"))]
             for p in raw["pages"])
    _xlsx_sheet(workbook, "Pages", PAGE_COLUMNS,
                {"source_file": 28, "warnings": 30, "text_layer_score": 16,
                 "ocr_confidence": 15, "needs_review": 13, "text": 100}, pages)
    blocks = ([cell_text(r[c]) if c in ("source_file", "text") else r[c] for c in BLOCK_COLUMNS]
              for r in iter_blocks(raw))
    _xlsx_sheet(workbook, "Blocks", BLOCK_COLUMNS,
                {"source_file": 28, "text": 90, "block_index": 12}, blocks)
    tmp = path.with_name(path.name + ".tmp")
    workbook.save(tmp)
    os.replace(tmp, path)


# ---------------------------------------------------------------- docx

def export_docx(raw: dict, path: Path) -> None:
    from docx import Document

    document = Document()
    name = source_name(raw)
    document.core_properties.title = name
    document.add_heading(name, level=1)
    pages = raw["pages"]
    for position, page in enumerate(pages):
        document.add_heading(f"Page {page['page']}", level=2)
        text = clean(page.get("text")).strip()
        for paragraph_text in re.split(r"\n\s*\n", text) if text else []:
            paragraph = document.add_paragraph()
            lines = paragraph_text.split("\n")
            for number, line in enumerate(lines):
                run = paragraph.add_run(line)
                if number < len(lines) - 1:
                    run.add_break()
        if position < len(pages) - 1:
            document.add_page_break()
    tmp = path.with_name(path.name + ".tmp")
    document.save(str(tmp))
    os.replace(tmp, path)


# ---------------------------------------------------------------- sqlite

_SQLITE_SCHEMA = """
CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY,
    file_name TEXT NOT NULL UNIQUE,
    source_file TEXT,
    sha256 TEXT,
    size_bytes INTEGER,
    page_count INTEGER,
    extraction_method TEXT,
    text_layer_score REAL,
    ocr_used INTEGER,
    schema_version TEXT,
    tool_version TEXT
);
CREATE TABLE IF NOT EXISTS pages (
    id INTEGER PRIMARY KEY,
    document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    source_file TEXT,
    page INTEGER NOT NULL,
    width REAL,
    height REAL,
    rotation INTEGER,
    method TEXT,
    text_layer_score REAL,
    confidence REAL,
    needs_review INTEGER,
    warnings TEXT,
    text TEXT,
    UNIQUE (document_id, page)
);
CREATE TABLE IF NOT EXISTS blocks (
    id INTEGER PRIMARY KEY,
    document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    page_id INTEGER NOT NULL REFERENCES pages(id) ON DELETE CASCADE,
    source_file TEXT,
    page INTEGER NOT NULL,
    block_index INTEGER NOT NULL,
    text TEXT,
    x1 REAL, y1 REAL, x2 REAL, y2 REAL,
    method TEXT,
    confidence REAL,
    UNIQUE (page_id, block_index)
);
CREATE INDEX IF NOT EXISTS blocks_document_page ON blocks (document_id, page);
"""


def export_sqlite(raw: dict, path: Path) -> None:
    """Idempotent: the document's previous rows are replaced inside one transaction."""
    name = source_name(raw)
    source = raw.get("source") or {}
    conn = sqlite3.connect(str(path))
    try:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.executescript(_SQLITE_SCHEMA)
        with conn:
            conn.execute("DELETE FROM documents WHERE file_name = ?", (name,))
            document_id = conn.execute(
                "INSERT INTO documents (file_name, source_file, sha256, size_bytes, page_count,"
                " extraction_method, text_layer_score, ocr_used, schema_version, tool_version)"
                " VALUES (?,?,?,?,?,?,?,?,?,?)",
                (name, raw.get("source_file"), source.get("sha256"), source.get("size_bytes"),
                 raw.get("page_count"), raw.get("extraction_method"), raw.get("text_layer_score"),
                 int(bool(raw.get("ocr_used"))), raw.get("schema_version"),
                 (raw.get("tool") or {}).get("version"))).lastrowid
            for page in raw["pages"]:
                page_id = conn.execute(
                    "INSERT INTO pages (document_id, source_file, page, width, height, rotation,"
                    " method, text_layer_score, confidence, needs_review, warnings, text)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (document_id, name, page["page"], page.get("width"), page.get("height"),
                     page.get("rotation"), page.get("method"), page_score(page),
                     page.get("ocr_confidence"), int(bool(page.get("needs_review"))),
                     json.dumps(page.get("warnings") or [], ensure_ascii=False),
                     page.get("text") or "")).lastrowid
                rows = []
                for index, block in enumerate(page.get("blocks") or []):
                    bbox = list(block.get("bbox") or [])[:4] + [None] * 4
                    rows.append((document_id, page_id, name, page["page"],
                                 block.get("order", index), block.get("text", ""),
                                 bbox[0], bbox[1], bbox[2], bbox[3],
                                 page.get("method"), block.get("confidence")))
                conn.executemany(
                    "INSERT INTO blocks (document_id, page_id, source_file, page, block_index,"
                    " text, x1, y1, x2, y2, method, confidence) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    rows)
    finally:
        conn.close()


EXPORTERS = {"html": export_html, "txt": export_txt, "md": export_md, "csv": export_csv,
             "xlsx": export_xlsx, "docx": export_docx, "jsonl": export_jsonl,
             "sqlite": export_sqlite}


# ---------------------------------------------------------------- CLI

def build_arg_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        description="Export <name>.raw.json to another format (extracts only if raw is missing).")
    ap.add_argument("--format", required=True, choices=sorted(FORMATS))
    ap.add_argument("--input", required=True, help="PDF file or folder containing PDFs")
    ap.add_argument("--output", required=True, help="folder with / for the .raw.json files")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--force-ocr", action="store_true",
                      help="only for documents that still need extraction")
    mode.add_argument("--native-only", action="store_true",
                      help="only for documents that still need extraction")
    ap.add_argument("--config", help="config JSON for the extraction step")
    return ap


def extract_missing(targets: list[Path], out_dir: Path, args) -> str:
    """Run the normal pipeline for PDFs without a raw.json. Returns its output."""
    extra = []
    if args.force_ocr:
        extra.append("--force-ocr")
    if args.native_only:
        extra.append("--native-only")
    if args.config:
        extra += ["--config", args.config]
    captured = io.StringIO()
    with contextlib.redirect_stdout(captured):
        for target in targets:
            extract.main(["--input", str(target), "--output", str(out_dir), *extra])
    return captured.getvalue()


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    args = build_arg_parser().parse_args(argv)
    input_path = Path(args.input).resolve()
    out_dir = Path(args.output).resolve()

    pdfs = {p.stem: p for p in extract.collect_pdfs(input_path)} if input_path.exists() else {}
    raws = ({p.name[:-len(RAW_SUFFIX)]: p for p in out_dir.glob("*" + RAW_SUFFIX)}
            if out_dir.is_dir() else {})
    if input_path.is_file():
        stems = sorted(pdfs)
    else:
        stems = sorted(set(pdfs) | set(raws), key=str.lower)
    if not stems:
        print(f"ERROR: no PDF in {input_path} and no {RAW_SUFFIX} in {out_dir}")
        return 2

    missing = [stem for stem in stems if stem not in raws]
    extraction_log = ""
    if missing:
        out_dir.mkdir(parents=True, exist_ok=True)
        # nothing extracted yet -> one normal folder run; otherwise only the missing files,
        # so existing raw.json files are never re-extracted
        whole_folder = not raws and input_path.is_dir()
        targets = [input_path] if whole_folder else [pdfs[stem] for stem in missing]
        extraction_log = extract_missing(targets, out_dir, args)

    exporter, suffix = EXPORTERS[args.format], FORMATS[args.format]
    exported, extracted, stale, errors = 0, 0, [], []
    for stem in stems:
        raw_path = out_dir / (stem + RAW_SUFFIX)
        if not raw_path.is_file():
            errors.append(f"{stem}: extraction failed (see {stem}.summary.json)")
            continue
        try:
            raw = json.loads(raw_path.read_text(encoding="utf-8"))
            exporter(raw, out_dir / (stem + suffix))
            exported += 1
            extracted += stem in missing
            pdf = pdfs.get(stem)
            if pdf is not None and stem not in missing \
                    and (raw.get("source") or {}).get("size_bytes") != pdf.stat().st_size:
                stale.append(pdf.name)
        except Exception as exc:
            errors.append(f"{stem}: {type(exc).__name__}: {exc}")
            (out_dir / (stem + suffix + ".tmp")).unlink(missing_ok=True)

    print(f"FORMAT: {args.format}")
    print(f"FILES: {len(stems)}")
    print(f"EXPORTED: {exported}")
    print(f"RAW_REUSED: {exported - extracted}")
    print(f"EXTRACTED: {extracted}")
    if stale:
        print(f"STALE_RAW: {', '.join(stale)} (PDF changed; run pdfjson to refresh)")
    print(f"ERRORS: {len(errors)}")
    for error in errors[:10]:
        print(f"  {error}")
    if errors and extraction_log:
        print(extraction_log.rstrip())
    print(f"OUTPUT: {out_dir}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
