"""pdfstruct.docwriter - Word and HTML content for documents created from text.

Used by pdfstruct.create only. Both writers take the structured blocks of
pdfstruct.inputs (heading, paragraph, list_item, table_row, code, rule, with optional
inline `spans`) and return the finished file content; neither touches the disk or the
network. Images are never embedded and links are never followed.
"""
from __future__ import annotations

import html
import io
import zipfile

from .export import clean
from .inputs.text import LINK_SCHEMES

CODE_FONT = "Consolas"
CODE_STYLE = "Code"
LINK_COLOR = "0563C1"
ZIP_TIME = (1980, 1, 1, 0, 0, 0)  # fixed, so the same text gives the same .docx bytes
MAX_TABLE_COLUMNS = 63            # Word's limit; a wider table is written as text rows


def _spans(block: dict) -> list[dict]:
    return block.get("spans") or [{"text": block.get("text", "")}]


def _tables(blocks: list[dict], index: int, key: str) -> tuple[list[list], int]:
    """Rows of the table that starts at `index` -> (rows, index after the table)."""
    table, rows = blocks[index].get("table"), []
    while (index < len(blocks) and blocks[index].get("kind") == "table_row"
           and blocks[index].get("table") == table):
        block = blocks[index]
        rows.append(block.get(key) or [[{"text": cell}] for cell in block.get("cells") or []])
        index += 1
    width = max(len(row) for row in rows)
    return [row + [[]] * (width - len(row)) for row in rows], index


# ---------------------------------------------------------------- docx

def _mono(font) -> None:
    from docx.oxml.ns import qn

    font.name = CODE_FONT
    fonts = font._element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("w:cs", "w:eastAsia"):
        fonts.set(qn(script), CODE_FONT)


def _runs(paragraph, spans: list[dict], bold: bool = False) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    for span in spans:
        lines = clean(span["text"]).split("\n")
        for number, line in enumerate(lines):
            if line:
                run = paragraph.add_run(line)
                run.bold = True if bold or span.get("bold") else None
                run.italic = True if span.get("italic") else None
                if span.get("code"):
                    _mono(run.font)
                target = span.get("href") or ""
                if target.lower().startswith(LINK_SCHEMES):  # a relationship, never fetched
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("r:id"), paragraph.part.relate_to(
                        target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                    run.font.color.rgb = RGBColor.from_string(LINK_COLOR)
                    run.font.underline = True
                    paragraph._p.append(link)
                    link.append(run._r)
            if number < len(lines) - 1:
                paragraph.add_run().add_break()


def _restarted_list(document, style) -> int | None:
    """A numbering instance that starts at 1 again, so every ordered list counts for itself."""
    try:
        numbering = document.part.numbering_part.numbering_definitions._numbering
        shared = numbering.num_having_numId(style.element.pPr.numPr.numId.val)
        instance = numbering.add_num(shared.abstractNumId.val)
        instance.add_lvlOverride(ilvl=0).add_startOverride(1)
        return instance.numId
    except (AttributeError, KeyError):
        return None  # a template without that list definition: the style's own numbering


def _rule(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders, line = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for name, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        line.set(qn("w:" + name), value)
    borders.append(line)
    paragraph._p.get_or_add_pPr().append(borders)


def docx_bytes(blocks: list[dict], title: str) -> bytes:
    """A macro-free .docx: Word heading styles, real lists, real tables, a `Code` style."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Mm, Pt
    from docx.table import _Cell

    document = Document()
    document.core_properties.title = clean(title)[:255]  # the format allows no more
    document.core_properties.author = "PDFStruct"
    for section in document.sections:  # A4 like the PDF; the library's template is US Letter
        section.page_width, section.page_height = Mm(210), Mm(297)
    code = document.styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = document.styles["Normal"]
    _mono(code.font)
    code.font.size = Pt(9.5)
    code.paragraph_format.left_indent = Pt(12)

    style_ids: dict[str, str] = {}

    def paragraph_in(style: str):
        """A new paragraph in a named style. The style is looked up once: python-docx walks
        the whole style sheet on every `add_paragraph(style=...)`, which adds up to minutes
        for a document with tens of thousands of headings or list items."""
        if style not in style_ids:
            style_ids[style] = document.styles[style].style_id
        paragraph = document.add_paragraph()
        paragraph._p.style = style_ids[style]
        return paragraph

    index, open_list, list_id = 0, None, None
    while index < len(blocks):
        block = blocks[index]
        kind = block.get("kind") or "paragraph"
        if kind != "list_item":
            open_list = None
        if kind == "table_row":
            start = index
            rows, index = _tables(blocks, index, "cell_spans")
            if len(rows[0]) > MAX_TABLE_COLUMNS:  # Word cannot open it: keep every cell as text
                for block in blocks[start:index]:
                    document.add_paragraph(clean(block.get("text", "")))
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            # the XML rows directly: table.cell(r, c) rebuilds its cell list on every call
            for r, (row, xml_row) in enumerate(zip(rows, table._tbl.tr_lst)):
                for cell, xml_cell in zip(row, xml_row.tc_lst):
                    _runs(_Cell(xml_cell, table).paragraphs[0], cell, bold=r == 0)  # first row: header
            document.add_paragraph()
            continue
        if kind == "heading":
            _runs(paragraph_in(f"Heading {min(max(block.get('level') or 1, 1), 9)}"), _spans(block))
        elif kind == "list_item":
            ordered = bool(block.get("ordered"))
            level = min(max(block.get("level") or 1, 1), 3)
            style = "List Number" if ordered else "List Bullet"
            paragraph = paragraph_in(style if level == 1 else f"{style} {level}")
            _runs(paragraph, _spans(block))
            if level == 1:
                if ordered and open_list != "ordered":
                    list_id = _restarted_list(document, paragraph.style)
                open_list = "ordered" if ordered else "bullet"
                if ordered and list_id is not None:
                    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numbering.get_or_add_ilvl().val = 0
                    numbering.get_or_add_numId().val = list_id
        elif kind == "code":
            _runs(paragraph_in(CODE_STYLE), [{"text": block.get("text", "")}])
        elif kind == "rule":
            _rule(document.add_paragraph())
        elif block.get("text", "").strip():
            _runs(document.add_paragraph(), _spans(block))
        index += 1

    buffer = io.BytesIO()
    document.save(buffer)
    fixed = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            entry = zipfile.ZipInfo(item.filename, ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(entry, source.read(item.filename))
    return fixed.getvalue()


# ---------------------------------------------------------------- html

_HTML_CSS = """
:root{color-scheme:light dark}
body{margin:0;font:16px/1.6 "Segoe UI",system-ui,sans-serif}
main{max-width:820px;margin:0 auto;padding:24px 16px}
h1,h2,h3,h4,h5,h6{line-height:1.25;margin:1.4em 0 .5em}
p{white-space:pre-wrap}
p,li{overflow-wrap:anywhere}
table{border-collapse:collapse;margin:1em 0}
th,td{border:1px solid #9a9a9a;padding:4px 10px;text-align:left;vertical-align:top}
pre{padding:10px 12px;border:1px solid #9a9a9a;border-radius:4px;overflow-x:auto}
code,pre{font-family:Consolas,"Cascadia Mono",monospace;font-size:.92em}
hr{border:0;border-top:1px solid #9a9a9a;margin:1.6em 0}
"""
# Nothing is loaded and nothing runs, whatever the text contains.
_HTML_POLICY = "default-src 'none'; style-src 'unsafe-inline'"


def _inline_html(spans: list[dict]) -> str:
    out = []
    for span in spans:
        piece = html.escape(clean(span["text"])).replace("\n", "<br>")
        if span.get("code"):
            piece = f"<code>{piece}</code>"
        if span.get("italic"):
            piece = f"<em>{piece}</em>"
        if span.get("bold"):
            piece = f"<strong>{piece}</strong>"
        target = span.get("href") or ""
        if target.lower().startswith(LINK_SCHEMES):
            piece = f'<a href="{html.escape(target, quote=True)}" rel="noopener noreferrer">{piece}</a>'
        out.append(piece)
    return "".join(out)


def html_text(blocks: list[dict], title: str) -> str:
    """One self-contained page. All text is escaped; no script, no remote resource."""
    out = ["<!doctype html>", '<html lang="und"><head><meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width,initial-scale=1">',
           f'<meta http-equiv="Content-Security-Policy" content="{_HTML_POLICY}">',
           f"<title>{html.escape(clean(title))}</title>", f"<style>{_HTML_CSS}</style></head>",
           "<body><main>"]
    lists: list[str] = []  # open list tags, outermost first

    def close_lists(depth: int = 0) -> None:
        while len(lists) > depth:
            out.append(f"</li></{lists.pop()}>")

    index = 0
    while index < len(blocks):
        block = blocks[index]
        kind = block.get("kind") or "paragraph"
        if kind == "list_item":
            tag = "ol" if block.get("ordered") else "ul"
            level = min(max(block.get("level") or 1, 1), len(lists) + 1)
            close_lists(level)
            if len(lists) == level and lists[-1] != tag:
                close_lists(level - 1)
            if len(lists) == level:
                out.append("</li>")
            else:
                out.append(f"<{tag}>")
                lists.append(tag)
            out.append(f"<li>{_inline_html(_spans(block))}")
            index += 1
            continue
        close_lists()
        if kind == "table_row":
            rows, index = _tables(blocks, index, "cell_spans")
            out.append("<table>")
            for number, row in enumerate(rows):
                cell = "th" if number == 0 else "td"
                out.append("<tr>" + "".join(f"<{cell}>{_inline_html(spans)}</{cell}>"
                                            for spans in row) + "</tr>")
            out.append("</table>")
            continue
        if kind == "heading":
            level = min(max(block.get("level") or 1, 1), 6)
            out.append(f"<h{level}>{_inline_html(_spans(block))}</h{level}>")
        elif kind == "code":
            out.append(f"<pre><code>{html.escape(clean(block.get('text', '')))}</code></pre>")
        elif kind == "rule":
            out.append("<hr>")
        elif block.get("text", "").strip():
            out.append(f"<p>{_inline_html(_spans(block))}</p>")
        index += 1
    close_lists()
    out.append("</main></body></html>")
    return "\n".join(out) + "\n"
