"""Documents created from text held in memory: the core service."""
import os
import zipfile

import pypdfium2 as pdfium
import pytest
from docx import Document

from pdfstruct import create, docwriter
from pdfstruct.inputs import text as text_input

PLACEHOLDERS = ["{{PROJECT_NAME}}", "{{REPO_PATH}}", "{{PRODUCTION_URL}}", "{{TARGET_USERS}}",
                "{{_PRIVATE_}}", "{{ spaced name }}"]
UNICODE = "Çağrı Öğretmen İstanbul ığüşöç – Straße größer Ärger"
MARKDOWN = f"""# Audit: {{{{PROJECT_NAME}}}}

Intro with **bold**, *italic*, `inline code` and {{{{REPO_PATH}}}} next to my_var_name.
{UNICODE}

## Steps

1. First {{{{PRODUCTION_URL}}}}
2. Second

Between the lists.

1. Restarted
2. Again

- Bullet one
  - Nested bullet
- Bullet {{{{_PRIVATE_}}}} and {{{{ spaced name }}}}

| Field | Value |
|-------|-------|
| Users | {{{{TARGET_USERS}}}} |
| **Strong** | a \\| b |

---

```bash
git clone {{{{REPO_PATH}}}}
  indented   spaces
```

See [the docs](https://example.com/a?b=1&c=2) or <https://example.org>.
"""
TEXT = f"First line {{{{PROJECT_NAME}}}}\n  indented second line\n\n# not a heading\n- not a list **x**\n\n{UNICODE}\n"


def make(tmp_path, name="DOC", content=MARKDOWN, formats=("docx",), **options):
    return create.create_document(create.CreateRequest(name=name, content=content,
                                                       formats=formats, output_dir=tmp_path,
                                                       **options))


def pdf_text(path) -> str:
    document = pdfium.PdfDocument(str(path))
    try:
        return "\n".join(page.get_textpage().get_text_range() for page in document)
    finally:
        document.close()


def list_id(paragraph):
    properties = paragraph._p.pPr
    return None if properties is None or properties.numPr is None else properties.numPr.numId.val


# ---------------------------------------------------------------- formats and names

def test_one_call_writes_every_format_with_the_exact_name(tmp_path):
    name = "WEBSITE_STRATEGY_AUDIT_TEMPLATE"
    result = make(tmp_path, name, formats=["docx", "pdf", "md", "html", "txt"])
    assert result.to_dict() == {"status": "created", "warnings": [], "outputs": [
        str(tmp_path / f"{name}.{ext}") for ext in ("docx", "pdf", "md", "html", "txt")]}
    assert sorted(p.name for p in tmp_path.iterdir()) == sorted(
        f"{name}.{ext}" for ext in ("docx", "pdf", "md", "html", "txt"))  # nothing else, no temp file
    assert (tmp_path / f"{name}.pdf").read_bytes().startswith(b"%PDF-")
    assert (tmp_path / f"{name}.md").read_text(encoding="utf-8") == MARKDOWN  # Markdown as written
    for ext in ("md", "html", "txt"):
        written = (tmp_path / f"{name}.{ext}").read_text(encoding="utf-8")
        assert UNICODE in written
        for placeholder in PLACEHOLDERS:
            assert placeholder in written, (ext, placeholder)


def test_default_format_is_docx_and_an_extension_in_the_name_is_not_doubled(tmp_path):
    assert create.CreateRequest("A", "x").formats == ("docx",)
    assert make(tmp_path, "Report.docx").outputs == [str(tmp_path / "Report.docx")]
    assert make(tmp_path, "notes.MD", formats=["md", "md", "TXT", ".txt"]).outputs == [
        str(tmp_path / "notes.md"), str(tmp_path / "notes.txt")]  # repeats are written once
    other = make(tmp_path, "plan.pdf", formats=["docx"])
    assert other.outputs == [str(tmp_path / "plan.docx")]
    assert other.warnings[0].startswith("name_extension_ignored")
    assert make(tmp_path, "v1.2 final", formats="md, txt").outputs == [
        str(tmp_path / "v1.2 final.md"), str(tmp_path / "v1.2 final.txt")]
    assert not list(tmp_path.glob("*.docx.docx")) and not list(tmp_path.glob("*.md.md"))


@pytest.mark.parametrize("name", [
    "", "   ", "..", ".", ".hidden", "..\\evil", "../evil", "a/b", "a\\b", "C:\\Temp\\x", "C:x",
    "/etc/passwd", "\\\\server\\share\\x", "x:stream", "a\x00b", "a\nb", "a\tb", "a\x7fb", "a<b",
    'a"b', "a|b", "a?b", "a*b", "trailing.", "CON", "con", "PRN", "AUX", "NUL", "nul.docx", "COM1",
    "com1.txt", "LPT1", "LPT9.report", "CONIN$", ".docx", "x" * 121, None, 7])
def test_unsafe_names_are_rejected_and_nothing_is_written(tmp_path, name):
    result = make(tmp_path / "out", name)
    assert result.status == "invalid" and result.error_code == "invalid_name" and result.outputs == []
    assert not (tmp_path / "out").exists() and list(tmp_path.iterdir()) == []


def test_bad_requests_are_rejected_before_any_work(tmp_path):
    out = tmp_path / "out"
    cases = [(dict(content=""), "empty_content"), (dict(content=" \n\t "), "empty_content"),
             (dict(content=None), "empty_content"), (dict(formats=["exe"]), "unsupported_format"),
             (dict(formats=["docx", "xlsx"]), "unsupported_format"), (dict(formats=[]), "unsupported_format"),
             (dict(content_type="html"), "invalid_content_type"),
             (dict(content="x" * (create.MAX_CONTENT_CHARS + 1)), "content_too_large")]
    for options, code in cases:
        result = make(out, **options)
        assert (result.status, result.error_code) == ("invalid", code), options
        assert len(result.error) <= 200
    assert not out.exists()
    blocker = tmp_path / "file.txt"
    blocker.write_text("x", encoding="utf-8")
    assert make(blocker).error_code == "invalid_output_dir"
    assert make(tmp_path, content="x" * create.MAX_CONTENT_CHARS, formats=["txt"]).ok


def test_without_a_folder_documents_go_to_the_users_documents(tmp_path, monkeypatch):
    monkeypatch.setattr(create.Path, "home", classmethod(lambda cls: tmp_path))
    assert create.default_output_dir() == tmp_path / "PDFStruct"  # no Documents folder here
    (tmp_path / "Documents").mkdir()
    assert create.default_output_dir() == tmp_path / "Documents" / "PDFStruct"
    result = create.create_document(create.CreateRequest("A", "text", ["txt"]))
    assert result.outputs == [str(tmp_path / "Documents" / "PDFStruct" / "A.txt")]


# ---------------------------------------------------------------- overwrite and temporary files

def test_existing_files_are_kept_unless_overwrite_is_set(tmp_path):
    assert make(tmp_path, "A", "first", ["md", "txt"]).ok
    conflict = make(tmp_path, "A", "second", ["docx", "txt", "md"])
    assert conflict.to_dict() == {
        "status": "conflict", "outputs": [], "warnings": [], "error_code": "already_exists",
        "error": "already exists, nothing was written: A.txt, A.md. Choose another name or "
                 "folder, or allow overwriting"}
    assert (tmp_path / "A.md").read_text(encoding="utf-8") == "first\n"
    assert not (tmp_path / "A.docx").exists()  # not even the format that was free
    assert make(tmp_path, "a.MD", "second", ["md"]).status == ("conflict" if os.name == "nt" else "created")
    replaced = make(tmp_path, "A", "second", ["md"], overwrite=True)
    assert replaced.ok and (tmp_path / "A.md").read_text(encoding="utf-8") == "second\n"
    assert (tmp_path / "A.txt").read_text(encoding="utf-8") == "first\n"  # only what was asked for


def test_a_file_that_appears_late_is_not_replaced(tmp_path, monkeypatch):
    target = tmp_path / "A.md"
    target.write_text("someone else", encoding="utf-8")
    with pytest.raises(FileExistsError):
        create._publish(target, b"mine", overwrite=False)
    monkeypatch.setattr(os, "link", lambda *args: (_ for _ in ()).throw(OSError("no hard links")))
    with pytest.raises(FileExistsError):
        create._publish(target, b"mine", overwrite=False)
    create._publish(tmp_path / "B.md", b"mine", overwrite=False)  # the fallback still writes
    assert target.read_text(encoding="utf-8") == "someone else"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["A.md", "B.md"]


def test_a_failing_writer_leaves_nothing_behind_and_does_not_quote_the_text(tmp_path, monkeypatch):
    def broken(blocks, title):
        raise RuntimeError("cannot render: " + blocks[0]["text"])
    monkeypatch.setattr(docwriter, "html_text", broken)
    result = make(tmp_path / "out", "A", "TOP SECRET LINE", ["md", "html", "txt"])
    assert result.to_dict() == {"status": "failed", "outputs": [], "warnings": [],
                                "error_code": "render_failed",
                                "error": "html could not be created (RuntimeError)"}
    assert not (tmp_path / "out").exists()

    def no_space(target, data, overwrite):
        raise OSError("disk full")
    monkeypatch.undo()
    monkeypatch.setattr(create, "_publish", no_space)
    failed = make(tmp_path, "B", "TOP SECRET LINE", ["md"])
    assert failed.status == "failed" and failed.error_code == "write_failed"
    assert "SECRET" not in str(failed.to_dict()) and list(tmp_path.iterdir()) == []


def test_no_temporary_file_survives_a_failed_write(tmp_path, monkeypatch):
    monkeypatch.setattr(os, "replace", lambda *args: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(PermissionError):
        create._publish(tmp_path / "A.md", b"data", overwrite=True)
    assert list(tmp_path.iterdir()) == []


def test_same_text_gives_the_same_bytes(tmp_path):
    for folder in ("a", "b"):
        assert make(tmp_path / folder, formats=["docx", "pdf", "html", "md", "txt"]).ok
    for ext in ("docx", "pdf", "html", "md", "txt"):
        assert (tmp_path / "a" / f"DOC.{ext}").read_bytes() == (tmp_path / "b" / f"DOC.{ext}").read_bytes(), ext


# ---------------------------------------------------------------- docx

def test_markdown_docx_has_real_word_structure(tmp_path):
    assert make(tmp_path).ok
    document = Document(str(tmp_path / "DOC.docx"))
    assert document.core_properties.title == "Audit: {{PROJECT_NAME}}"
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert paragraphs == [
        ("Heading 1", "Audit: {{PROJECT_NAME}}"),
        ("Normal", "Intro with bold, italic, inline code and {{REPO_PATH}} next to my_var_name. " + UNICODE),
        ("Heading 2", "Steps"),
        ("List Number", "First {{PRODUCTION_URL}}"), ("List Number", "Second"),
        ("Normal", "Between the lists."),
        ("List Number", "Restarted"), ("List Number", "Again"),
        ("List Bullet", "Bullet one"), ("List Bullet 2", "Nested bullet"),
        ("List Bullet", "Bullet {{_PRIVATE_}} and {{ spaced name }}"),
        ("Code", "git clone {{REPO_PATH}}\n  indented   spaces"),
        ("Normal", "See the docs or https://example.org.")]

    by_text = {p.text: p for p in document.paragraphs}
    first, second = list_id(by_text["First {{PRODUCTION_URL}}"]), list_id(by_text["Restarted"])
    assert first and second and first != second                  # the second list counts from 1 again
    assert list_id(by_text["Second"]) == first and list_id(by_text["Again"]) == second

    intro = by_text[paragraphs[1][1]]
    styled = {run.text: (run.bold, run.italic, run.font.name) for run in intro.runs}
    assert styled["bold"] == (True, None, None) and styled["italic"] == (None, True, None)
    assert styled["inline code"] == (None, None, "Consolas")
    assert document.styles["Code"].font.name == "Consolas"

    (table,) = document.tables
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["Field", "Value"], ["Users", "{{TARGET_USERS}}"], ["Strong", "a | b"]]
    assert table.cell(0, 0).paragraphs[0].runs[0].bold and table.cell(2, 0).paragraphs[0].runs[0].bold
    assert not table.cell(1, 0).paragraphs[0].runs[0].bold

    links = {rel.target_ref for rel in document.part.rels.values() if rel.is_external}
    assert links == {"https://example.com/a?b=1&c=2", "https://example.org"}
    xml = document.element.xml
    assert xml.count("<w:hyperlink ") == 2 and "the docs" in xml and "<w:pBdr>" in xml  # link text, rule


def test_plain_text_docx_keeps_paragraphs_line_breaks_and_literal_characters(tmp_path):
    assert make(tmp_path, content=TEXT, content_type="text").ok
    document = Document(str(tmp_path / "DOC.docx"))
    assert [(p.style.name, p.text) for p in document.paragraphs] == [
        ("Normal", "First line {{PROJECT_NAME}}\n  indented second line"),
        ("Normal", "# not a heading\n- not a list **x**"), ("Normal", UNICODE)]
    assert document.core_properties.title == "DOC" and not document.tables


def test_docx_is_macro_free_and_embeds_nothing_remote(tmp_path):
    content = MARKDOWN + "\n![logo](https://example.com/logo.png) and <img src='https://example.com/x.png'>\n"
    result = make(tmp_path, content=content)
    assert result.warnings == ["images_not_embedded: pictures are not loaded; their alt text is kept"]
    with zipfile.ZipFile(tmp_path / "DOC.docx") as archive:
        names = archive.namelist()
        types = archive.read("[Content_Types].xml").decode()
        body = archive.read("word/document.xml").decode()
    assert not [n for n in names if "vba" in n.lower() or n.startswith("word/media") or n.endswith(".bin")]
    assert "macroEnabled" not in types and "<w:drawing" not in body and "<w:pict" not in body
    assert "logo" in body and "&lt;img src" in body  # kept as text


# ---------------------------------------------------------------- pdf / html / txt / md

def test_markdown_pdf_is_a_readable_reflow_with_all_text(tmp_path):
    result = make(tmp_path, formats=["pdf"])
    assert result.ok and result.warnings == []  # every character is in the PDF font
    text = " ".join(pdf_text(tmp_path / "DOC.pdf").split())
    for expected in ("Audit: {{PROJECT_NAME}}", UNICODE, "1. First {{PRODUCTION_URL}}", "2. Second",
                     "1. Restarted", "• Bullet one", "{{TARGET_USERS}}", "git clone {{REPO_PATH}}",
                     "the docs", "my_var_name"):
        assert expected in text, expected
    document = pdfium.PdfDocument(str(tmp_path / "DOC.pdf"))
    assert [round(v) for v in document[0].get_size()] == [595, 842]  # A4
    document.close()
    odd = make(tmp_path, "ODD", "Zeichen 漢字 ✓", ["pdf"])
    assert odd.ok and odd.warnings[0].startswith("pdf_missing_glyphs")


def test_html_escapes_everything_and_can_load_or_run_nothing(tmp_path):
    hostile = ("# T <script>alert(1)</script>\n\n<img src=x onerror=alert(2)> <iframe src='//evil'></iframe>\n\n"
               "[click](javascript:alert(3)) [data](data:text/html;base64,AAAA) [ok](https://example.com/?a=\"b\")\n\n"
               "![pic](https://example.com/p.png) `</code><script>x</script>`\n\n"
               "```\n</pre><script>alert(4)</script>\n```\n\n| a | <b onclick=x>b</b> |\n|---|---|\n| 1 | 2 |\n")
    assert make(tmp_path, "</title><script>x", hostile, ["html"]).status == "invalid"  # not a file name
    assert make(tmp_path, content=hostile, formats=["html"]).ok
    page = (tmp_path / "DOC.html").read_text(encoding="utf-8")
    lowered = page.lower()
    assert "<script" not in lowered and "<img" not in lowered and "<iframe" not in lowered
    assert " onerror=" not in lowered.replace("&lt;img src=x onerror=", "") and "<b onclick" not in lowered
    assert "javascript:" not in page.split("<body>")[1].replace("(javascript:alert", "")  # never a link
    assert page.count("<a href=") == 1 and '<a href="https://example.com/?a=&quot;b&quot;"' in page
    assert "default-src 'none'" in page and "src=\"" not in page and "<link" not in lowered and "@import" not in page
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page and "&lt;/pre&gt;&lt;script&gt;" in page
    assert "<title>T &lt;script&gt;alert(1)&lt;/script&gt;</title>" in page

    assert make(tmp_path, "STRUCT", formats=["html"]).ok
    page = (tmp_path / "STRUCT.html").read_text(encoding="utf-8")
    for expected in ("<h1>Audit: {{PROJECT_NAME}}</h1>", "<h2>Steps</h2>", "<strong>bold</strong>",
                     "<em>italic</em>", "<code>inline code</code>", "<ol>\n<li>First {{PRODUCTION_URL}}",
                     "<ul>\n<li>Nested bullet", "<tr><th>Field</th><th>Value</th></tr>",
                     "<tr><td>Users</td><td>{{TARGET_USERS}}</td></tr>", "<hr>",
                     "<pre><code>git clone {{REPO_PATH}}\n  indented   spaces</code></pre>",
                     '<a href="https://example.com/a?b=1&amp;c=2" rel="noopener noreferrer">the docs</a>'):
        assert expected in page, expected
    assert page.count("<ol>") == page.count("</ol>") == 2 and page.count("<ul>") == page.count("</ul>") == 2
    assert page.count("<li>") == page.count("</li>") == 7


def test_markdown_txt_drops_the_markup_and_keeps_the_text(tmp_path):
    assert make(tmp_path, formats=["txt"]).ok
    assert (tmp_path / "DOC.txt").read_text(encoding="utf-8") == f"""Audit: {{{{PROJECT_NAME}}}}

Intro with bold, italic, inline code and {{{{REPO_PATH}}}} next to my_var_name. {UNICODE}

Steps

1. First {{{{PRODUCTION_URL}}}}
2. Second

Between the lists.

1. Restarted
2. Again
- Bullet one
  - Nested bullet
- Bullet {{{{_PRIVATE_}}}} and {{{{ spaced name }}}}

Field | Value
Users | {{{{TARGET_USERS}}}}
Strong | a | b

----------------------------------------

git clone {{{{REPO_PATH}}}}
  indented   spaces

See the docs (https://example.com/a?b=1&c=2) or https://example.org.
"""


def test_plain_text_is_never_read_as_markup(tmp_path):
    assert make(tmp_path, content=TEXT.replace("\n", "\r\n"), formats=["txt", "md", "html", "pdf"],
                content_type="text").ok
    assert (tmp_path / "DOC.txt").read_bytes() == TEXT.encode("utf-8")  # byte for byte, LF endings
    markdown = (tmp_path / "DOC.md").read_text(encoding="utf-8")
    assert "\\# not a heading" in markdown and "\\- not a list **x**" in markdown
    assert "First line {{PROJECT_NAME}}  \n" in markdown  # hard break keeps the line order
    page = (tmp_path / "DOC.html").read_text(encoding="utf-8")
    assert "<p>First line {{PROJECT_NAME}}<br>  indented second line</p>" in page
    assert "<h1>" not in page and "<li>" not in page and "<strong>" not in page
    assert "# not a heading" in pdf_text(tmp_path / "DOC.pdf")


def test_inline_markup_rules():
    def styled(text):
        return [(s["text"], "".join(k[0] for k in ("bold", "italic", "code") if s[k]) + ("L" if s["href"] else ""))
                for s in text_input.inline_spans(text)]
    assert styled("a **b *c* d** e") == [("a ", ""), ("b ", "b"), ("c", "bi"), (" d", "b"), (" e", "")]
    assert styled("***x*** __y__ _z_") == [("x", "bi"), (" ", ""), ("y", "b"), (" ", ""), ("z", "i")]
    assert styled("snake_case_name 2 * 3 * 4 a_b_c") == [("snake_case_name 2 * 3 * 4 a_b_c", "")]
    assert styled(r"\*not\* `**code**` \{\{x\}\}") == [("*", ""), ("not", ""), ("*", ""), (" ", ""),
                                                      ("**code**", "c"), (" ", ""), ("{", ""), ("{", ""),
                                                      ("x", ""), ("}", ""), ("}", "")]
    assert styled("{{A_B}} *{{_c_}}* {{**d**}}") == [("{{A_B}}", ""), (" ", ""), ("{{_c_}}", "i"),
                                                    (" ", ""), ("{{**d**}}", "")]
    assert styled("[t](ftp://x) [u](mailto:a@b.c)") == [("t", ""), (" (ftp://x)", ""), (" ", ""), ("u", "L")]
    assert text_input.spans_text(text_input.inline_spans("<https://a.b> [x](https://c.d)")) == \
        "https://a.b x (https://c.d)"


def test_only_a_header_with_a_matching_rule_line_is_a_table():
    text = "a | b\n---\n\nx | y\n--|:-:\n1 | 2\nlast | row\n\nprice | tax\nno rule here\n"
    blocks = text_input.markdown_blocks(text_input.split_lines(text), rich=True)
    assert [(b["kind"], b["text"]) for b in blocks] == [
        ("heading", "a | b"), ("table_row", "x | y"), ("table_row", "1 | 2"), ("table_row", "last | row"),
        ("paragraph", "price | tax no rule here")]
    assert {b["table"] for b in blocks if b["kind"] == "table_row"} == {1}


def test_file_inputs_are_read_exactly_as_before():
    lines = text_input.split_lines(MARKDOWN)
    plain = text_input.markdown_blocks(lines)
    assert all("spans" not in block and block["kind"] != "rule" for block in plain)
    assert not any(block["kind"] == "table_row" for block in plain)
    assert "| Field | Value |\n|-------|-------|" in "\n".join(b["text"] for b in plain)  # literal
    assert plain[1]["text"].startswith("Intro with **bold**, *italic*, `inline code`") and "\n" in plain[1]["text"]
