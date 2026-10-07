"""Input adapters: DOCX, TXT, Markdown, HTML and images through the shared service."""
import csv
import json
import socket
from pathlib import Path

import jsonschema
import pytest

from conftest import (TOOL_DIR, UNICODE_LINE, make_docx, make_image, make_text_pdf, make_tiff)
from pdfstruct import cli, export, extract, inputs, ocr, service
from pdfstruct.service import JobRequest, run_job

SCHEMA = json.loads((TOOL_DIR / "schemas" / "raw_document.schema.json").read_text(encoding="utf-8"))


def convert(path, formats, **options):
    result = run_job(JobRequest(inputs=[path], formats=formats, **options))
    return result.files[0]


def raw_of(item) -> dict:
    raw = json.loads(Path(item.outputs["json"]).read_text(encoding="utf-8"))
    jsonschema.validate(raw, SCHEMA)
    return raw


def kinds(raw) -> list:
    return [(b["kind"], b.get("level")) for b in raw["pages"][0]["blocks"]]


# ---------------------------------------------------------------- TXT

def test_txt_is_unpaged_and_prints_no_page_headings(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_bytes(("﻿" + UNICODE_LINE + "\nzweite Zeile\n\n\nNeuer Absatz\n").encode("utf-8"))
    item = convert(source, ["json", "txt", "md", "html", "docx", "csv", "jsonl", "sqlite", "xlsx"])
    assert item.status == "succeeded" and item.input_format == "txt" and item.pages == 1
    assert item.method == "native" and item.ocr_used is False
    raw = raw_of(item)
    assert raw["schema_version"] == "1.1" and raw["input_format"] == "txt" and raw["paged"] is False
    assert Path(item.outputs["json"]).name == "notes.txt.raw.json"
    blocks = raw["pages"][0]["blocks"]
    assert [b["text"] for b in blocks] == [UNICODE_LINE + "\nzweite Zeile", "Neuer Absatz"]
    assert all("bbox" not in b and b["kind"] == "paragraph" for b in blocks)  # no invented coordinates
    assert raw["pages"][0]["width"] is None and raw["coordinate_system"]["unit"] is None

    out = tmp_path / "output"
    text = (out / "notes.txt.txt").read_text(encoding="utf-8")
    assert "Page" not in text and "=====" not in text and UNICODE_LINE in text
    markdown = (out / "notes.txt.md").read_text(encoding="utf-8")
    assert "## Page" not in markdown and "# notes.txt" not in markdown and "Neuer Absatz" in markdown
    page = (out / "notes.txt.html").read_text(encoding="utf-8")
    assert "<h2>Page 1</h2>" not in page and "<h2>Content</h2>" in page and "<nav>" not in page
    from docx import Document
    paragraphs = [p.text for p in Document(str(out / "notes.txt.docx")).paragraphs]
    assert not any(p.startswith("Page ") for p in paragraphs) and "Neuer Absatz" in paragraphs
    with (out / "notes.txt.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2 and rows[0]["x1"] == "" and rows[0]["method"] == "native"
    record = json.loads((out / "notes.txt.jsonl").read_text(encoding="utf-8").splitlines()[0])
    assert record["bbox"] is None
    assert source.read_bytes().startswith(b"\xef\xbb\xbf")  # the source is untouched


def test_txt_encoding_fallback_is_reported(tmp_path):
    source = tmp_path / "alt.txt"
    source.write_bytes("Größe über Maß".encode("cp1252"))
    item = convert(source, ["json"])
    assert any(w.startswith("encoding_fallback") and "cp1252" in w for w in item.warnings)
    assert raw_of(item)["pages"][0]["text"] == "Größe über Maß"
    utf16 = tmp_path / "wide.txt"
    utf16.write_bytes("Çağrı".encode("utf-16"))
    wide = convert(utf16, ["json"])
    assert not wide.warnings and raw_of(wide)["pages"][0]["text"] == "Çağrı"


def test_empty_txt(tmp_path):
    source = tmp_path / "leer.txt"
    source.write_bytes(b"")
    item = convert(source, ["json", "txt", "md"])
    assert item.status == "succeeded" and item.method == "none"
    assert (tmp_path / "output" / "leer.txt.txt").read_text(encoding="utf-8") == ""


def test_ocr_options_on_text_input_are_reported_not_ignored(tmp_path):
    source = tmp_path / "a.txt"
    source.write_text("Hallo", encoding="utf-8")
    forced = convert(source, ["json"], mode="force-ocr")
    assert forced.status == "succeeded" and forced.ocr_used is False
    assert any(w.startswith("ocr_mode_ignored") and "force-ocr" in w for w in forced.warnings)
    native = convert(source, ["json"], mode="native-only")
    assert any(w.startswith("ocr_mode_ignored") for w in native.warnings)


# ---------------------------------------------------------------- Markdown

MARKDOWN = """# Titel

Erster Absatz
mit zwei Zeilen.

## Liste

- Punkt A
  - Unterpunkt
1. Nummer eins

```python
print("<b>kein HTML</b>")
```

| a | b |
|---|---|
| 1 | 2 |

<script>alert(1)</script>

Setext
------
"""


def test_markdown_structure_without_invented_semantics(tmp_path):
    source = tmp_path / "doc.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    item = convert(source, ["json", "md", "txt"])
    raw = raw_of(item)
    assert raw["input_format"] == "md"
    assert kinds(raw) == [("heading", 1), ("paragraph", None), ("heading", 2), ("list_item", 1),
                          ("list_item", 2), ("list_item", 1), ("code", None), ("paragraph", None),
                          ("paragraph", None), ("heading", 2)]
    blocks = raw["pages"][0]["blocks"]
    assert blocks[1]["text"] == "Erster Absatz\nmit zwei Zeilen."
    assert blocks[5]["ordered"] is True and blocks[3]["ordered"] is False
    assert blocks[6]["text"] == 'print("<b>kein HTML</b>")'
    assert blocks[7]["text"].startswith("| a | b |")       # a table stays literal text
    assert blocks[8]["text"] == "<script>alert(1)</script>"  # embedded HTML is text, never run
    markdown = (tmp_path / "output" / "doc.md.md").read_text(encoding="utf-8")
    assert markdown.startswith("# Titel\n") and "## Liste" in markdown and "  - Unterpunkt" in markdown
    assert "```\nprint(" in markdown
    assert any(w.startswith("same_format_roundtrip") for w in item.warnings)


# ---------------------------------------------------------------- HTML

HTML = """<!doctype html><html><head><title>Nicht sichtbar</title>
<style>p{color:red}</style><script>window.geheim = "SCRIPTTEXT"</script></head>
<body><h1>Bericht &amp; Analyse</h1>
<p>Erster <b>Absatz</b><br>zweite Zeile</p>
<div hidden><p>VERSTECKT <span>tief</p></div>
<p style="display:none">UNSICHTBAR</p>
<ul><li>Eins<ul><li>Eins A</li></ul></li><li>Zwei</ol>
<ol><li>Schritt</li></ol>
<table><tr><th>Name</th><th>Preis</th></tr><tr><td>Apfel</td><td>1,20 &euro;</td></tr></table>
<img src="https://example.invalid/x.png" alt="ALTTEXT">
<pre>  code
    eingerueckt</pre>
<p>Schluss</p></body></html>"""


def test_html_visible_text_only_and_no_network(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("the HTML adapter must not open a network connection")
    monkeypatch.setattr(socket, "socket", no_network)
    monkeypatch.setattr(socket, "create_connection", no_network)
    source = tmp_path / "seite.html"
    source.write_text(HTML, encoding="utf-8")
    item = convert(source, ["json", "csv", "html"])
    raw = raw_of(item)
    text = raw["pages"][0]["text"]
    for hidden in ("SCRIPTTEXT", "VERSTECKT", "UNSICHTBAR", "color:red", "Nicht sichtbar", "ALTTEXT"):
        assert hidden not in text
    assert kinds(raw) == [("heading", 1), ("paragraph", None), ("list_item", 1), ("list_item", 2),
                          ("list_item", 1), ("list_item", 1), ("table_row", None),
                          ("table_row", None), ("code", None), ("paragraph", None)]
    blocks = raw["pages"][0]["blocks"]
    assert blocks[0]["text"] == "Bericht & Analyse"
    assert blocks[1]["text"] == "Erster Absatz\nzweite Zeile"
    assert blocks[5]["ordered"] is True and blocks[2]["ordered"] is False
    assert blocks[7]["cells"] == ["Apfel", "1,20 €"] and blocks[7]["table"] == 0
    assert blocks[8]["text"] == "  code\n    eingerueckt"
    assert blocks[9]["text"] == "Schluss"  # content after the hidden block is not swallowed
    assert any(w.startswith("tables_flattened") for w in item.warnings)       # csv is not a table
    assert any(w.startswith("same_format_roundtrip") for w in item.warnings)  # html to html


def test_html_declared_charset(tmp_path):
    source = tmp_path / "tr.htm"
    source.write_bytes('<html><head><meta charset="iso-8859-9"></head><body><p>Çağrı şöyle</p></body></html>'
                       .encode("iso-8859-9"))
    item = convert(source, ["json"])
    assert raw_of(item)["pages"][0]["text"] == "Çağrı şöyle" and item.input_format == "html"


# ---------------------------------------------------------------- DOCX

def test_docx_structure_in_body_order(tmp_path):
    source = make_docx(tmp_path / "bericht.docx")
    item = convert(source, ["json", "md", "docx", "xlsx"])
    raw = raw_of(item)
    assert raw["input_format"] == "docx" and raw["paged"] is False and raw["page_count"] == 1
    assert kinds(raw) == [("heading", 1), ("paragraph", None), ("heading", 2), ("list_item", 1),
                          ("list_item", 1), ("list_item", 1), ("table_row", None),
                          ("table_row", None), ("paragraph", None)]
    blocks = raw["pages"][0]["blocks"]
    assert blocks[1]["text"] == UNICODE_LINE and blocks[5]["ordered"] is True
    assert blocks[6]["cells"] == ["Name", "Menge", "Preis"] and blocks[7]["cells"] == ["Apfel", "3", "1,20"]
    assert any(w.startswith("docx_layout_not_preserved") for w in item.warnings)
    assert any(w.startswith("same_format_roundtrip") for w in item.warnings)
    assert any(w.startswith("tables_flattened") for w in item.warnings)
    assert len(item.warnings) <= service.MAX_WARNINGS
    markdown = (tmp_path / "output" / "bericht.docx.md").read_text(encoding="utf-8")
    assert markdown.startswith("# Quartalsbericht\n") and "| Name | Menge | Preis |" in markdown
    assert "| --- | --- | --- |" in markdown and "1. Schritt eins" in markdown
    from docx import Document
    again = Document(str(tmp_path / "output" / "bericht.docx.docx"))
    assert again.tables[0].cell(1, 0).text == "Apfel"
    assert [p.style.name for p in again.paragraphs if p.text == "Quartalsbericht"] == ["Heading 1"]


def test_invalid_docx_fails_with_a_short_actionable_error(tmp_path):
    source = tmp_path / "kaputt.docx"
    source.write_bytes(b"this is not a zip archive")
    item = convert(source, ["json"])
    assert item.status == "failed" and item.error_code == "invalid_docx"
    assert "not a readable .docx" in item.error and len(item.error) <= service.MAX_MESSAGE_CHARS


# ---------------------------------------------------------------- images

def test_native_only_on_an_image_is_an_explicit_failure(tmp_path):
    source = make_image(tmp_path / "foto.png")
    item = convert(source, ["json"], mode="native-only")
    assert item.status == "failed" and item.error_code == "native_only_unsupported"
    assert "no text layer" in item.error
    assert not list((tmp_path / "output").glob("*.raw.json"))


def test_image_without_ocr_fails_clearly(tmp_path, monkeypatch):
    monkeypatch.setattr(ocr, "available", lambda: False)
    item = convert(make_image(tmp_path / "foto.png"), ["json", "txt"])
    assert item.status == "failed" and item.error_code == "ocr_unavailable"
    assert "OCR support" in item.error and item.outputs == {}


def test_unreadable_image(tmp_path):
    source = tmp_path / "kein.png"
    source.write_bytes(b"not an image")
    item = convert(source, ["json"])
    assert item.status == "failed" and item.error_code == "invalid_image"


def test_images_go_straight_to_ocr(tmp_path, ocr_available, monkeypatch):
    if not ocr_available:
        pytest.skip("OCR packages not installed on this platform")
    import pypdfium2
    monkeypatch.setattr(pypdfium2, "PdfDocument",  # no intermediate PDF for OCR
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("no PDF expected")))
    make_image(tmp_path / "scan.png")
    make_image(tmp_path / "scan.bmp")
    make_image(tmp_path / "scan.webp", lossless=True)
    events = []
    result = run_job(JobRequest(inputs=sorted(tmp_path.glob("scan.*")), formats=["json", "txt", "docx"]),
                     on_event=events.append)
    assert [f.status for f in result.files] == ["succeeded"] * 3
    assert [f.input_format for f in result.files] == ["bmp", "png", "webp"]
    assert result.ocr_engine["engine"].startswith("paddleocr")
    for item in result.files:
        raw = raw_of(item)
        page = raw["pages"][0]
        assert raw["schema_version"] == "1.1" and raw["paged"] is True and raw["ocr_used"] is True
        assert raw["coordinate_system"]["unit"] == "px" and (page["width"], page["height"]) == (1240, 700)
        assert page["method"] == "ocr" and "Hallo" in page["text"] and "12345" in page["text"]
        for block in page["blocks"]:
            x1, y1, x2, y2 = block["bbox"]
            assert 0 <= x1 < x2 <= 1241 and 0 <= y1 < y2 <= 701 and block["kind"] == "line"
            assert 0 < block["confidence"] <= 1
        assert any(w.startswith("ocr_text_only") for w in item.warnings)  # docx holds text only
    assert sorted(p.name for p in (tmp_path / "output").glob("*.raw.json")) == [
        "scan.bmp.raw.json", "scan.png.raw.json", "scan.webp.raw.json"]
    assert sum(1 for e in events if e == {"type": "phase", "phase": "ocr", "page": 1}) == 3


def test_exif_orientation_and_multipage_tiff(tmp_path, ocr_available):
    if not ocr_available:
        pytest.skip("OCR packages not installed on this platform")
    from PIL import Image
    upright = make_image(tmp_path / "upright.png")
    with Image.open(upright) as image:
        exif = Image.Exif()
        exif[0x0112] = 6  # stored rotated; a viewer turns it 90 degrees clockwise
        image.rotate(90, expand=True).save(tmp_path / "gedreht.jpg", exif=exif, quality=95)
    item = convert(tmp_path / "gedreht.jpg", ["json"])
    page = raw_of(item)["pages"][0]
    assert (page["width"], page["height"]) == (1240, 700)  # upright again
    assert "Hallo" in page["text"] and item.input_format == "jpeg"

    tiff = convert(make_tiff(tmp_path / "mehr.tiff"), ["json", "txt"])
    raw = raw_of(tiff)
    assert tiff.pages == 2 and raw["page_count"] == 2 and tiff.input_format == "tiff"
    assert "Alpha" in raw["pages"][0]["text"] and "Beta" in raw["pages"][1]["text"]
    text = (tmp_path / "output" / "mehr.tiff.txt").read_text(encoding="utf-8")
    assert "===== Page 2 / 2 =====" in text  # frames are real pages


# ---------------------------------------------------------------- names, collisions, sources

def test_same_stem_different_types_never_collide(tmp_path):
    make_text_pdf(tmp_path / "report.pdf")
    make_docx(tmp_path / "report.docx")
    (tmp_path / "report.txt").write_text("Nur Text", encoding="utf-8")
    result = run_job(JobRequest(inputs=sorted(tmp_path.glob("report.*")), formats=["json", "txt", "html"]))
    assert result.ok and [f.input_format for f in result.files] == ["docx", "pdf", "txt"]
    names = sorted(p.name for p in (tmp_path / "output").iterdir() if not p.name.endswith(".summary.json"))
    assert names == ["report.docx.html", "report.docx.raw.json", "report.docx.txt",
                     "report.html", "report.raw.json", "report.txt",       # the PDF keeps v0.1.0 names
                     "report.txt.html", "report.txt.raw.json", "report.txt.txt"]
    pdf_raw = json.loads((tmp_path / "output" / "report.raw.json").read_text(encoding="utf-8"))
    jsonschema.validate(pdf_raw, SCHEMA)
    assert pdf_raw["schema_version"] == "1.0" and "input_format" not in pdf_raw and "paged" not in pdf_raw
    assert "Beispieltext" in (tmp_path / "output" / "report.txt").read_text(encoding="utf-8")
    assert (tmp_path / "output" / "report.txt.txt").read_text(encoding="utf-8") == "Nur Text\n"
    assert (tmp_path / "report.txt").read_text(encoding="utf-8") == "Nur Text"


def test_an_input_is_never_overwritten(tmp_path):
    make_text_pdf(tmp_path / "a.pdf")
    (tmp_path / "a.html").write_text("<p>ORIGINAL</p>", encoding="utf-8")
    result = run_job(JobRequest(inputs=[tmp_path / "a.pdf", tmp_path / "a.html"],
                                formats=["html"], output_dir=tmp_path))  # output next to the sources
    pdf_item, html_item = result.files
    assert pdf_item.status == "failed" and pdf_item.error_code == "export_failed"
    assert "never overwritten" in pdf_item.error
    assert (tmp_path / "a.html").read_text(encoding="utf-8") == "<p>ORIGINAL</p>"
    assert html_item.status == "succeeded" and Path(html_item.outputs["html"]).name == "a.html.html"


def test_pdf_to_pdf_is_skipped_without_extraction(tmp_path, monkeypatch):
    make_text_pdf(tmp_path / "a.pdf")
    calls = []
    real = extract.process_pdf
    monkeypatch.setattr(extract, "process_pdf", lambda *a, **k: calls.append(1) or real(*a, **k))
    only = convert(tmp_path / "a.pdf", ["pdf"])
    assert only.status == "skipped" and only.skipped_formats == {"pdf": "already_pdf"}
    assert calls == [] and not (tmp_path / "output").exists() and only.outputs == {}
    both = convert(tmp_path / "a.pdf", ["txt", "pdf"])
    assert both.status == "succeeded" and list(both.outputs) == ["txt"]
    assert both.skipped_formats == {"pdf": "already_pdf"} and calls == [1]
    assert not (tmp_path / "output" / "a.pdf").exists()


# ---------------------------------------------------------------- CLI and inspect

def test_cli_accepts_the_new_inputs(tmp_path, capsys):
    (tmp_path / "notes.txt").write_text("Hallo CLI", encoding="utf-8")
    make_docx(tmp_path / "b.docx")
    make_text_pdf(tmp_path / "c.pdf")
    assert cli.run("json", ["notes.txt"], cwd=tmp_path) == 0          # alias with an explicit file
    assert (tmp_path / "output" / "notes.txt.raw.json").is_file()
    assert cli.run("md", ["b.docx"], cwd=tmp_path) == 0
    assert (tmp_path / "output" / "b.docx.md").is_file()
    capsys.readouterr()

    assert cli.run("txt", [], cwd=tmp_path) == 0                      # no argument: PDFs only, as before
    assert "FILES: 3" in capsys.readouterr().out                      # c.pdf + the two existing raws
    assert cli.main(["--format", "csv", "--all-types"], cwd=tmp_path) == 0
    assert (tmp_path / "output" / "b.docx.csv").is_file() and (tmp_path / "output" / "c.csv").is_file()

    assert cli.main(["b.docx", "--format", "pdf,txt"], cwd=tmp_path) == 0
    printed = capsys.readouterr().out
    assert "✓ PDF" in printed and (tmp_path / "output" / "b.docx.pdf").stat().st_size > 1000
    assert cli.main(["c.pdf", "--format", "pdf"], cwd=tmp_path) == 0  # single format: PDF to PDF
    printed = capsys.readouterr().out
    assert "SKIPPED: 1 (already PDF: c)" in printed and "EXPORTED: 0" in printed
    assert not (tmp_path / "output" / "c.pdf").exists()
    (tmp_path / "x.xyz").write_text("x", encoding="utf-8")
    assert cli.main(["x.xyz", "--format", "txt"], cwd=tmp_path) == 2
    assert "Desteklenmeyen dosya türü" in capsys.readouterr().out


def test_interactive_menu_lists_every_supported_type(tmp_path):
    import io
    (tmp_path / "a.md").write_text("# Hallo", encoding="utf-8")
    make_text_pdf(tmp_path / "b.pdf")
    (tmp_path / "ignored.xyz").write_text("x", encoding="utf-8")
    keys = iter(["a", "enter", "space", "up", "space", "enter"])  # all files; JSON off, PDF on
    out = io.StringIO()
    code = cli.main([], cwd=tmp_path, read_key=lambda: next(keys), write=out.write,
                    read_line=lambda: "")
    text = out.getvalue()
    assert code == 0 and "a.md" in text and "b.pdf" in text and "ignored.xyz" not in text
    assert "[ ] PDF" in text and "✓ PDF" in text and "skipped (already_pdf)" in text
    assert (tmp_path / "output" / "a.md.pdf").is_file() and not (tmp_path / "output" / "b.pdf").exists()


def test_inspect_returns_facts_but_no_text(tmp_path):
    make_text_pdf(tmp_path / "a.pdf", pages=4)
    (tmp_path / "n.txt").write_text("GEHEIMER INHALT", encoding="utf-8")
    make_tiff(tmp_path / "t.tif")
    pdf = service.inspect_file(tmp_path / "a.pdf")
    assert pdf["input_format"] == "pdf" and pdf["pages"] == 4 and pdf["estimated_method"] == "native"
    txt = service.inspect_file(tmp_path / "n.txt")
    assert txt["pages"] == 1 and txt["paged"] is False and "GEHEIM" not in json.dumps(txt)
    tif = service.inspect_file(tmp_path / "t.tif")
    assert tif["pages"] == 2 and tif["estimated_method"] == "ocr" and tif["input_format"] == "tiff"
    assert service.inspect_file(tmp_path / "nope.pdf")["error_code"] == "not_found"
    (tmp_path / "x.xyz").write_text("x", encoding="utf-8")
    assert service.inspect_file(tmp_path / "x.xyz")["error_code"] == "unsupported_input"
    assert not (tmp_path / "output").exists()  # inspecting writes nothing


def test_supported_extensions_are_the_documented_ones():
    assert set(inputs.EXTENSIONS) == {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm",
                                      ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
    assert service.OUTPUT_FORMATS == ["json", "html", "txt", "md", "csv", "xlsx", "docx", "jsonl",
                                      "sqlite", "pdf"]
    assert set(export.EXPORTERS) == set(export.FORMATS)
