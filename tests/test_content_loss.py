"""Text that readers used to drop: DOCX fields/insertions/nested tables/notes, HTML form values
and nested tables, PDF form values and free-text annotations."""
import json
import shutil
import zipfile
from pathlib import Path

from lxml import etree

from pdfstruct import extract, inputs
from pdfstruct.service import JobRequest, run_job

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def convert(path, formats=("json",)):
    return run_job(JobRequest(inputs=[path], formats=list(formats))).files[0]


def raw_of(item) -> dict:
    return json.loads(Path(item.outputs["json"]).read_text(encoding="utf-8"))


def blocks_of(item) -> list[dict]:
    return raw_of(item)["pages"][0]["blocks"]


# ---------------------------------------------------------------- DOCX

def make_rich_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_paragraph("Fließtext")
    pieces = {
        "Link: ": '<w:hyperlink w:anchor="x"><w:r><w:t>Verweis</w:t></w:r></w:hyperlink>',
        "Feld: ": '<w:sdt><w:sdtPr/><w:sdtContent><w:r><w:t>Inhaltssteuerelement</w:t></w:r></w:sdtContent></w:sdt>',
        "Neu: ": '<w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>eingefügt</w:t></w:r></w:ins>',
        "Alt: ": '<w:del w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>GELÖSCHT</w:delText></w:r></w:del>',
        "Datum: ": '<w:fldSimple w:instr=" DATE "><w:r><w:t>15.10.2026</w:t></w:r></w:fldSimple>',
        "Code: ": '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE </w:instrText></w:r>'
                  '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r>'
                  '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        "Tag: ": '<w:smartTag w:uri="u" w:element="e"><w:r><w:t>Ort</w:t></w:r></w:smartTag>',
    }
    for label, xml in pieces.items():
        paragraph = document.add_paragraph(label)
        wrapper = etree.fromstring(f'<w:wrap xmlns:w="{W}">{xml}</w:wrap>')
        for child in list(wrapper):
            paragraph._p.append(child)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Außen"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Innen"
    table.cell(0, 1).text = "Rechts"
    document.save(str(path))
    _add_notes_and_comment(path)
    return path


def _add_notes_and_comment(path: Path) -> None:
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    ctype = "application/vnd.openxmlformats-officedocument.wordprocessingml."
    parts = {
        "footnotes": f'<w:footnotes xmlns:w="{W}"><w:footnote w:type="separator" w:id="-1"><w:p><w:r>'
                     f'<w:separator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>'
                     f'Fußnotentext</w:t></w:r></w:p></w:footnote></w:footnotes>',
        "endnotes": f'<w:endnotes xmlns:w="{W}"><w:endnote w:id="1"><w:p><w:r><w:t>Endnotentext'
                    f'</w:t></w:r></w:p></w:endnote></w:endnotes>',
        "comments": f'<w:comments xmlns:w="{W}"><w:comment w:id="0" w:author="a"><w:p><w:r><w:t>'
                    f'KOMMENTAR</w:t></w:r></w:p></w:comment></w:comments>',
    }
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"Flie\xc3\x9ftext</w:t></w:r>",
                                    b'Flie\xc3\x9ftext</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>'
                                    b'<w:r><w:endnoteReference w:id="1"/></w:r>', 1)
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", "".join(
                    f'<Relationship Id="rId{n}x" Type="{rel}{n}" Target="{n}.xml"/>' for n in parts
                ).encode() + b"</Relationships>")
            elif item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", "".join(
                    f'<Override PartName="/word/{n}.xml" ContentType="{ctype}{n}+xml"/>' for n in parts
                ).encode() + b"</Types>")
            zout.writestr(item, data)
        for name, xml in parts.items():
            zout.writestr(f"word/{name}.xml", xml)
    shutil.move(tmp, path)


def test_docx_reads_text_inside_links_fields_controls_and_insertions(tmp_path):
    item = convert(make_rich_docx(tmp_path / "reich.docx"))
    texts = [b["text"] for b in blocks_of(item)]
    for expected in ("Link: Verweis", "Feld: Inhaltssteuerelement", "Neu: eingefügt", "Alt:",
                     "Datum: 15.10.2026", "Code: 7", "Tag: Ort"):
        assert expected in texts, expected
    whole = raw_of(item)["pages"][0]["text"]
    assert "GELÖSCHT" not in whole and "PAGE" not in whole  # deletions and field codes stay out
    assert any(w.startswith("docx_tracked_changes") for w in item.warnings)


def test_docx_nested_table_text_stays_in_its_cell(tmp_path):
    item = convert(make_rich_docx(tmp_path / "reich.docx"))
    (row,) = [b for b in blocks_of(item) if b["kind"] == "table_row"]
    assert row["cells"] == ["Außen Innen", "Rechts"]
    assert any(w.startswith("docx_nested_tables_flattened") for w in item.warnings)


def test_docx_footnotes_and_endnotes_are_read_and_comments_only_counted(tmp_path):
    item = convert(make_rich_docx(tmp_path / "reich.docx"), ["json", "md"])
    regions = {b["text"]: b.get("region") for b in blocks_of(item)}
    assert regions["Fußnotentext"] == "footnote" and regions["Endnotentext"] == "endnote"
    assert "KOMMENTAR" not in raw_of(item)["pages"][0]["text"]
    assert "docx_comments_not_read: 1 comment(s) are not document text" in item.warnings
    markdown = Path(item.outputs["md"]).read_text(encoding="utf-8")
    assert "Fußnotentext" in markdown


def test_docx_with_a_broken_notes_part_keeps_its_body(tmp_path):
    path = make_rich_docx(tmp_path / "kaputt.docx")
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            zout.writestr(item, b"<w:footnotes" if item.filename == "word/footnotes.xml" else data)
    shutil.move(tmp, path)
    item = convert(path)
    assert item.status == "succeeded" and "Fließtext" in raw_of(item)["pages"][0]["text"]
    assert any(w.startswith("docx_footnotes_unreadable") for w in item.warnings)
    assert "Endnotentext" in raw_of(item)["pages"][0]["text"]  # the other notes part still works


# ---------------------------------------------------------------- HTML

HTML = """<!doctype html><html><head><title>TITEL</title></head><body>
<table><tr><td>Außen<table><tr><td>Innen</td><td>Zwei</td></tr></table></td><td>Rechts</td></tr>
<tr><td>Zeile zwei</td></tr></table>
<form><p>Name <input type="text" value="Mustafa"> Passwort <input type="password" value="GEHEIM">
<input type="hidden" value="VERSTECKT"></p>
<textarea>Freitext im Formular</textarea>
<p>Land <select><option>Schweiz</option><option selected>Deutschland</option></select>
Stadt <select><option>Dinslaken</option><option>Duisburg</option></select></p>
<datalist><option>LISTE</option></datalist><button>KNOPF</button></form>
<p><img src="x.png" alt="ALTTEXT">Ende</p></body></html>"""


def test_html_keeps_nested_table_text_and_the_outer_cells(tmp_path):
    source = tmp_path / "seite.html"
    source.write_text(HTML, encoding="utf-8")
    item = convert(source)
    rows = [b["cells"] for b in blocks_of(item) if b["kind"] == "table_row"]
    assert rows == [["Außen Innen Zwei", "Rechts"], ["Zeile zwei"]]
    assert any(w.startswith("html_nested_tables_flattened") for w in item.warnings)


def test_html_reads_visible_form_values_only(tmp_path):
    source = tmp_path / "seite.html"
    source.write_text(HTML, encoding="utf-8")
    item = convert(source)
    texts = [b["text"] for b in blocks_of(item)]
    assert "Name Mustafa Passwort" in texts
    assert "Freitext im Formular" in texts
    assert "Land Deutschland Stadt Dinslaken" in texts  # selected option, else the first one
    whole = raw_of(item)["pages"][0]["text"]
    for hidden in ("GEHEIM", "VERSTECKT", "LISTE", "KNOPF", "TITEL", "ALTTEXT", "Duisburg", "Schweiz"):
        assert hidden not in whole, hidden
    assert "Ende" in texts


# ---------------------------------------------------------------- PDF

def make_form_pdf(path: Path, value="Mustafa Kelesoglu", with_note=True) -> Path:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.drawString(72, 760, "Antrag auf Zeugnisbewertung")
    c.drawString(72, 720, "Name:")
    c.acroForm.textfield(name="name", value=value, x=140, y=712, width=220, height=20)
    c.acroForm.checkbox(name="einverstanden", tooltip="Einverstanden", x=140, y=680, checked=True)
    c.acroForm.checkbox(name="abgelehnt", x=180, y=680, checked=False)
    c.acroForm.textfield(name="leer", value="", x=140, y=640, width=220, height=20)
    if with_note:
        c.drawString(72, 600, "Unterschrift")
    c.save()
    return path


def test_pdf_form_values_are_read_with_field_name_and_box(tmp_path):
    item = convert(make_form_pdf(tmp_path / "antrag.pdf"), ["json", "txt"])
    page = raw_of(item)["pages"][0]
    forms = [b for b in page["blocks"] if b.get("region") == "form"]
    assert [(b["field"], b["text"]) for b in forms] == [("name", "Mustafa Kelesoglu"),
                                                         ("einverstanden", "[x] Einverstanden")]
    assert all(len(b["bbox"]) == 4 for b in forms)
    assert "Mustafa Kelesoglu" in page["text"] and page["method"] == "native"
    assert "Mustafa Kelesoglu" in Path(item.outputs["txt"]).read_text(encoding="utf-8")


def test_pdf_free_text_annotation_is_read_and_notes_are_counted(tmp_path):
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    from reportlab.pdfgen import canvas
    source = tmp_path / "notiz.pdf"
    c = canvas.Canvas(str(source))
    c.drawString(72, 760, "Seitentext")
    c.save()
    document = pdfium.PdfDocument(str(source))
    page = document[0]
    for subtype, content in ((pdfium_c.FPDF_ANNOT_FREETEXT, "Sichtbarer Hinweis"),
                             (pdfium_c.FPDF_ANNOT_TEXT, "Haftnotiz")):
        annot = pdfium_c.FPDFPage_CreateAnnot(page.raw, subtype)
        rect = pdfium_c.FS_RECTF(72, 700, 300, 680)
        pdfium_c.FPDFAnnot_SetRect(annot, rect)
        value = (content + "\0").encode("utf-16-le")
        pdfium_c.FPDFAnnot_SetStringValue(annot, b"Contents",
                                          extract.ctypes.cast(extract.ctypes.create_string_buffer(value, len(value)),
                                                              extract.ctypes.POINTER(pdfium_c.FPDF_WCHAR)))
        pdfium_c.FPDFPage_CloseAnnot(annot)
    page.close()
    patched = tmp_path / "notiz2.pdf"
    document.save(str(patched))
    document.close()
    page = raw_of(convert(patched))["pages"][0]
    assert "Sichtbarer Hinweis" in page["text"] and "Haftnotiz" not in page["text"]
    assert [b["region"] for b in page["blocks"] if "region" in b] == ["annotation"]
    assert "pdf_comments_not_read: 1 comment annotation(s) are not page text" in page["warnings"]


def test_only_pdfs_with_a_form_get_a_new_fingerprint(tmp_path):
    from reportlab.pdfgen import canvas
    plain = tmp_path / "plain.pdf"
    c = canvas.Canvas(str(plain))
    c.drawString(72, 760, "Nur Text")
    c.save()
    assert extract._pdf_revision(plain) == ""  # cached plain PDFs (and their OCR) stay valid
    assert extract._pdf_revision(make_form_pdf(tmp_path / "form.pdf")) == extract.FORMS_REVISION


def test_reader_revisions_refresh_docx_and_html_caches():
    assert inputs.READER_REVISIONS["docx"] == "3" and inputs.READER_REVISIONS["html"] == "2"
    assert "pdf" not in inputs.READER_REVISIONS
