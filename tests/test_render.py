"""DOCX text boxes / headers / footers in the raw document, and layout PDF through an office engine."""
import subprocess
from pathlib import Path

import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
import pytest
from lxml import etree

from conftest import make_docx
from pdfstruct import export, inputs, mcp_server, render, service
from pdfstruct.service import JobRequest, run_job

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
      'xmlns:v="urn:schemas-microsoft-com:vml"')


def convert(path, formats=("json", "pdf"), **options):
    return run_job(JobRequest(inputs=[path], formats=list(formats), **options)).files[0]


def _box_paragraphs(lines) -> str:
    return "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in lines)


def _text_box_run(lines, doc_pr_id=50, with_fallback=False) -> etree._Element:
    """A floating text box run as Word writes it (optionally with the VML fallback copy)."""
    drawing = f"""<w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
      relativeHeight="2" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="page"><wp:posOffset>180000</wp:posOffset></wp:positionH>
      <wp:positionV relativeFrom="page"><wp:posOffset>1800000</wp:posOffset></wp:positionV>
      <wp:extent cx="1800000" cy="1800000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
      <wp:docPr id="{doc_pr_id}" name="Sidebar"/><wp:cNvGraphicFramePr/>
      <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1800000" cy="1800000"/>
      </a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></wps:spPr>
      <wps:txbx><w:txbxContent>{_box_paragraphs(lines)}</w:txbxContent></wps:txbx>
      <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"/>
      </wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>"""
    if with_fallback:
        inner = (f'<mc:AlternateContent><mc:Choice Requires="wps">{drawing}</mc:Choice><mc:Fallback>'
                 f'<w:pict><v:rect style="width:140pt;height:140pt"><v:textbox><w:txbxContent>'
                 f'{_box_paragraphs(lines)}</w:txbxContent></v:textbox></v:rect></w:pict>'
                 f'</mc:Fallback></mc:AlternateContent>')
    else:
        inner = drawing
    return etree.fromstring(f"<w:r {NS}>{inner}</w:r>")


def make_layout_docx(path: Path, with_fallback=False, image: Path | None = None) -> Path:
    """Header, footer, a body with a floating sidebar text box (a CV, in short)."""
    from docx import Document
    from docx.shared import Cm

    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "KOPFZEILE Bewerbung"
    section.footer.paragraphs[0].text = "FUSSZEILE Seite"
    anchor = document.add_paragraph("Hauptteil Anfang")
    anchor._p.append(_text_box_run(["SIDEBAR Kontakt", "mail@example.org"], with_fallback=with_fallback))
    document.add_paragraph("Hauptteil Ende")
    if image is not None:
        document.add_picture(str(image), width=Cm(3))
    document.save(str(path))
    return path


def _png(path: Path) -> Path:
    from PIL import Image
    Image.new("RGB", (60, 60), (27, 100, 164)).save(path)
    return path


def _pdf_text_and_images(path) -> tuple[int, str, int]:
    document = pdfium.PdfDocument(str(path))
    try:
        texts, images = [], 0
        for page in document:
            textpage = page.get_textpage()
            texts.append(textpage.get_text_range())
            images += sum(1 for obj in page.get_objects() if obj.type == pdfium_c.FPDF_PAGEOBJ_IMAGE)
            textpage.close()
            page.close()
        return len(document), " ".join(" ".join(texts).split()), images
    finally:
        document.close()


def _fake_engine(engine="word"):
    """render_pdf stand-in: writes a real one-page PDF and names the engine."""
    def fake(source, target, *args, **kwargs):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(target))
        c.drawString(72, 720, "LAYOUT")
        c.save()
        return engine
    return fake


# ---------------------------------------------------------------- reader

@pytest.mark.parametrize("with_fallback", [False, True])
def test_reader_reads_header_text_box_and_footer_once_in_reading_order(tmp_path, with_fallback):
    source = make_layout_docx(tmp_path / "cv.docx", with_fallback=with_fallback)
    blocks = inputs.read(source, "docx")["pages"][0]["blocks"]
    assert [(b["text"], b.get("region")) for b in blocks] == [
        ("KOPFZEILE Bewerbung", "header"),
        ("Hauptteil Anfang", None),
        ("SIDEBAR Kontakt", "textbox"),          # right after its anchor paragraph
        ("mail@example.org", "textbox"),         # the VML fallback copy is not read twice
        ("Hauptteil Ende", None),
        ("FUSSZEILE Seite", "footer"),
    ]
    assert [b["order"] for b in blocks] == list(range(len(blocks)))


def test_plain_docx_blocks_carry_no_region(tmp_path):
    blocks = inputs.read(make_docx(tmp_path / "bericht.docx"), "docx")["pages"][0]["blocks"]
    assert blocks and not any("region" in b for b in blocks)


def test_images_in_body_are_counted(tmp_path):
    source = make_layout_docx(tmp_path / "cv.docx", image=_png(tmp_path / "foto.png"))
    warnings = inputs.read(source, "docx")["warnings"]
    assert "docx_images_skipped: 1 embedded image(s) not extracted" in warnings


def test_cached_docx_raw_is_refreshed_when_the_reader_revision_changes(tmp_path, monkeypatch):
    source = make_layout_docx(tmp_path / "cv.docx")
    assert convert(source, ["json"]).raw_reused is False
    assert convert(source, ["json"]).raw_reused is True
    monkeypatch.setitem(inputs.READER_REVISIONS, "docx", "next")
    assert convert(source, ["json"]).raw_reused is False
    assert "pdf" not in inputs.READER_REVISIONS  # PDF fingerprints (OCR caches) stay valid


# ---------------------------------------------------------------- pdf_mode in the service

def test_auto_uses_the_layout_engine_and_other_formats_still_come_from_raw(tmp_path, monkeypatch):
    monkeypatch.setattr(render, "render_pdf", _fake_engine("word"))
    item = convert(make_layout_docx(tmp_path / "cv.docx"), ["json", "md", "pdf"])
    assert item.status == "succeeded" and item.pdf_renderer == "word"
    _, text, _ = _pdf_text_and_images(item.outputs["pdf"])
    assert text == "LAYOUT"
    markdown = Path(item.outputs["md"]).read_text(encoding="utf-8")
    assert "SIDEBAR Kontakt" in markdown and "KOPFZEILE Bewerbung" in markdown
    assert not any(w.startswith("pdf_layout") for w in item.warnings)


def test_auto_falls_back_to_reflow_with_a_warning_when_no_engine_exists(tmp_path, monkeypatch):
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "none")
    item = convert(make_layout_docx(tmp_path / "cv.docx"))
    assert item.status == "succeeded" and item.pdf_renderer == "reflow"
    assert any(w.startswith("pdf_layout_unavailable: no layout renderer found") for w in item.warnings)
    _, text, _ = _pdf_text_and_images(item.outputs["pdf"])
    assert "SIDEBAR Kontakt" in text  # the re-flow now carries text box content too


def test_auto_falls_back_when_the_engine_fails(tmp_path, monkeypatch):
    def broken(*args, **kwargs):
        raise render.RenderFailed("word: COMError boom")
    monkeypatch.setattr(render, "render_pdf", broken)
    item = convert(make_layout_docx(tmp_path / "cv.docx"))
    assert item.status == "succeeded" and item.pdf_renderer == "reflow"
    assert any(w.startswith("pdf_layout_failed: word: COMError boom") for w in item.warnings)


def test_layout_mode_fails_instead_of_falling_back(tmp_path, monkeypatch):
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "none")
    item = convert(make_layout_docx(tmp_path / "cv.docx"), pdf_mode="layout")
    assert item.status == "failed" and item.error_code == "export_failed"
    assert "no layout renderer found" in item.error and "json" in item.outputs
    assert not (tmp_path / "output" / "cv.docx.pdf").exists()
    assert not list((tmp_path / "output").glob("*.tmp"))


def test_reflow_mode_never_starts_an_engine(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("render_pdf must not run in reflow mode")
    monkeypatch.setattr(render, "render_pdf", forbidden)
    item = convert(make_layout_docx(tmp_path / "cv.docx"), pdf_mode="reflow")
    assert item.status == "succeeded" and item.pdf_renderer == "reflow"
    assert not any(w.startswith("pdf_layout") for w in item.warnings)  # reflow was asked for


def test_text_inputs_are_never_rendered(tmp_path, monkeypatch):
    monkeypatch.setattr(render, "render_pdf", _fake_engine("word"))
    source = tmp_path / "notiz.md"
    source.write_text("# Titel\n\nText", encoding="utf-8")
    item = convert(source, ["pdf"])
    assert item.pdf_renderer == "reflow"


def test_unknown_pdf_mode_is_rejected_everywhere(tmp_path):
    with pytest.raises(ValueError, match="unknown pdf mode"):
        run_job(JobRequest(inputs=[tmp_path / "x.docx"], formats=["pdf"], pdf_mode="pretty"))
    answer = mcp_server.convert([str(tmp_path / "x.docx")], ["pdf"], pdf_mode="pretty")
    assert answer["error_code"] == "bad_request" and "pdf_mode" in answer["error"]


def test_mcp_reports_the_engine_only_when_one_was_used(tmp_path, monkeypatch):
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "none")  # inherited by the child process
    answer = mcp_server.convert([str(make_layout_docx(tmp_path / "cv.docx"))], ["pdf"])
    (entry,) = answer["files"]
    assert entry["status"] == "succeeded" and "pdf_renderer" not in entry
    assert any(w.startswith("pdf_layout_unavailable") for w in entry["warnings"])


def test_service_child_accepts_pdf_mode(tmp_path):
    request = tmp_path / "request.json"
    result = tmp_path / "result.json"
    source = make_layout_docx(tmp_path / "cv.docx")
    request.write_text(service.json.dumps({"inputs": [str(source)], "formats": ["pdf"],
                                           "pdf_mode": "reflow"}), encoding="utf-8")
    assert service.main([str(request), str(result)]) == 0
    files = service.json.loads(result.read_text(encoding="utf-8"))["files"]
    assert files[0]["pdf_renderer"] == "reflow"


# ---------------------------------------------------------------- export command

def test_export_command_renders_docx_and_reports_it(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(render, "render_pdf", _fake_engine("libreoffice"))
    source = make_layout_docx(tmp_path / "cv.docx")
    assert export.main(["--format", "pdf", "--input", str(source), "--output", str(tmp_path / "out")]) == 0
    printed = capsys.readouterr().out
    assert "LAYOUT_RENDERED: cv.docx (libreoffice)" in printed and "EXPORTED: 1" in printed
    assert export.main(["--format", "pdf", "--input", str(source), "--output", str(tmp_path / "out"),
                        "--pdf-mode", "reflow"]) == 0
    assert "LAYOUT_RENDERED" not in capsys.readouterr().out


# ---------------------------------------------------------------- render module

def test_render_pdf_tries_engines_in_order(tmp_path, monkeypatch):
    calls = []

    def attempt(engine, source, target, timeout):
        calls.append(engine)
        if engine == "word":
            raise render.RenderFailed("Word did not finish")
        target.write_bytes(b"%PDF-1.4")

    monkeypatch.setattr(render, "_render_with", attempt)
    assert render.render_pdf(tmp_path / "a.docx", tmp_path / "a.pdf", engines=["word", "libreoffice"]) \
        == "libreoffice"
    assert calls == ["word", "libreoffice"]

    def always(engine, source, target, timeout):
        raise render.RenderFailed(f"{engine} broke")
    monkeypatch.setattr(render, "_render_with", always)
    with pytest.raises(render.RenderFailed, match="word: word broke; libreoffice: libreoffice broke"):
        render.render_pdf(tmp_path / "a.docx", tmp_path / "a.pdf", engines=["word", "libreoffice"])
    with pytest.raises(render.RenderUnavailable):
        render.render_pdf(tmp_path / "a.docx", tmp_path / "a.pdf", engines=[])


def test_engine_choice_follows_the_environment(monkeypatch):
    monkeypatch.setattr(render, "_word_installed", lambda: True)
    monkeypatch.setattr(render, "soffice_path", lambda: "/usr/bin/soffice")
    monkeypatch.delenv("PDFSTRUCT_RENDERER", raising=False)
    assert render.available_engines() == ["word", "libreoffice"]
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "libreoffice")
    assert render.available_engines() == ["libreoffice"]
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "none")
    assert render.available_engines() == []
    monkeypatch.setattr(render, "soffice_path", lambda: None)
    monkeypatch.setenv("PDFSTRUCT_RENDERER", "libreoffice")
    assert render.available_engines() == []


def test_a_hanging_word_is_stopped_and_only_its_own_process_is_killed(tmp_path, monkeypatch):
    (tmp_path / "word.pid").write_text("4242 4343")
    killed = []

    def hang(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd="word", timeout=1)

    monkeypatch.setattr(render.subprocess, "run", hang)
    monkeypatch.setattr(render, "_kill", lambda pids: killed.extend(pids))
    with pytest.raises(render.RenderFailed, match="did not finish within 1 s"):
        render._run_word(tmp_path / "input.docx", tmp_path / "input.pdf", tmp_path, timeout=1)
    assert killed == [4242, 4343]


def test_an_unusable_result_is_rejected_and_nothing_is_left_behind(tmp_path, monkeypatch):
    def fake_word(copy, produced, work, timeout):
        produced.write_bytes(b"not a pdf")

    monkeypatch.setattr(render, "_run_word", fake_word)
    source = make_docx(tmp_path / "a.docx")
    with pytest.raises(render.RenderFailed, match="does not open"):
        render.render_pdf(source, tmp_path / "a.pdf", engines=["word"])
    assert sorted(p.name for p in tmp_path.iterdir()) == ["a.docx"]


# ---------------------------------------------------------------- real engines (skipped when absent)

def _word_pids():
    return render._word_pids() if render._word_installed() else set()


@pytest.mark.skipif(not render._word_installed(), reason="Microsoft Word is not installed")
def test_word_renders_text_box_header_and_image_without_touching_the_source(tmp_path):
    folder = tmp_path / "Bewerbung Ünlü"  # a space and non-ASCII letters in the path
    folder.mkdir()
    source = make_layout_docx(folder / "Lebenslauf Çağrı.docx", image=_png(tmp_path / "foto.png"))
    before_files = sorted(p.name for p in folder.iterdir())
    before_word = _word_pids()
    target = tmp_path / "out.pdf"
    assert render.render_pdf(source, target, engines=["word"]) == "word"
    pages, text, images = _pdf_text_and_images(target)
    assert pages == 1 and images >= 1
    for expected in ("KOPFZEILE Bewerbung", "SIDEBAR Kontakt", "mail@example.org", "Hauptteil Ende",
                     "FUSSZEILE Seite"):
        assert expected in text, expected
    assert sorted(p.name for p in folder.iterdir()) == before_files  # no lock or temp file
    assert _word_pids() <= before_word                                 # our Word has quit


@pytest.mark.skipif(render.soffice_path() is None, reason="LibreOffice is not installed")
def test_libreoffice_renders_text_box_and_header(tmp_path):
    source = make_layout_docx(tmp_path / "cv.docx")
    assert render.render_pdf(source, tmp_path / "out.pdf", engines=["libreoffice"]) == "libreoffice"
    _, text, _ = _pdf_text_and_images(tmp_path / "out.pdf")
    assert "SIDEBAR Kontakt" in text and "KOPFZEILE Bewerbung" in text
