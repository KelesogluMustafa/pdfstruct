"""Local MCP server: content-free conversions, capped excerpts, real stdio protocol."""
import asyncio
import importlib.util
import json
import sys
from pathlib import Path

import pytest

from conftest import TOOL_DIR, make_docx, make_text_pdf
from pdfstruct import mcp_server

SECRET = "STRENG GEHEIMER VERTRAGSINHALT 4711"


@pytest.fixture
def doc(tmp_path) -> Path:
    path = tmp_path / "vertrag.txt"
    path.write_text(SECRET + "\n\n" + "Absatz mit Fülltext. " * 400, encoding="utf-8")
    return path


def test_supported_formats_lists_everything_briefly():
    info = mcp_server.supported_formats()
    assert info["outputs"] == ["json", "html", "txt", "md", "csv", "xlsx", "docx", "jsonl", "sqlite", "pdf"]
    assert ".docx" in info["inputs"]["docx"] and ".webp" in info["inputs"]["images"]
    assert len(mcp_server._dump(info)) < 1500


def test_inspect_has_no_document_text(doc, tmp_path):
    info = mcp_server.inspect(str(doc))
    assert info["ok"] and info["input_format"] == "txt" and info["pages"] == 1
    assert "GEHEIM" not in json.dumps(info)
    assert mcp_server.inspect(str(tmp_path / "nope.pdf"))["error_code"] == "not_found"


def test_convert_runs_in_a_child_process_and_returns_names_only(doc, tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("the server process must not run the conversion itself")
    monkeypatch.setattr(mcp_server.service, "run_job", forbidden)
    make_text_pdf(tmp_path / "a.pdf", pages=2)
    answer = mcp_server.convert([str(doc), str(tmp_path / "a.pdf"), str(tmp_path / "nope.docx")],
                                ["txt", "json", "pdf"])
    text = mcp_server._dump(answer)
    assert "GEHEIM" not in text and "Fülltext" not in text and "Beispieltext" not in text
    assert len(text) < 2000  # roughly 500 tokens for a normal answer
    assert (answer["succeeded"], answer["failed"], answer["ok"]) == (2, 1, False)
    first, second, third = answer["files"]
    assert first.pop("warnings")[0].startswith("same_format_roundtrip")  # txt to txt is flagged
    assert first == {"file": "vertrag.txt", "status": "succeeded", "type": "txt", "method": "native",
                     "pages": 1, "output_dir": str(tmp_path / "output"),
                     "outputs": {"json": "vertrag.txt.raw.json", "txt": "vertrag.txt.txt",
                                 "pdf": "vertrag.txt.pdf"}}
    assert second["outputs"] == {"json": "a.raw.json", "txt": "a.txt"}
    assert second["skipped"] == {"pdf": "already_pdf"} and second["pages"] == 2
    assert third["error_code"] == "not_found" and "outputs" not in third
    assert (tmp_path / "output" / "vertrag.txt.pdf").is_file()
    assert doc.read_text(encoding="utf-8").startswith(SECRET)  # source untouched


def test_convert_rejects_bad_requests_without_starting_work(tmp_path):
    assert mcp_server.convert([], ["txt"])["error_code"] == "bad_request"
    bad = mcp_server.convert([str(tmp_path / "a.pdf")], ["exe"])
    assert bad["error_code"] == "bad_request" and "unknown format" in bad["error"]
    assert not (tmp_path / "output").exists()


def test_many_files_keep_the_answer_small(tmp_path):
    paths = []
    for n in range(mcp_server.MAX_FILES_IN_DETAIL + 3):
        path = tmp_path / f"notiz_{n:02d}.txt"
        path.write_text(f"Inhalt {n}", encoding="utf-8")
        paths.append(str(path))
    answer = mcp_server.convert(paths + [str(tmp_path / "fehlt.txt")], ["txt"])
    assert answer["succeeded"] == len(paths) and answer["failed"] == 1
    assert [f["file"] for f in answer["files"]] == ["fehlt.txt"]  # only what needs attention
    assert answer["files_omitted"] == len(paths) and answer["output_dirs"] == [str(tmp_path / "output")]
    assert len(mcp_server._dump(answer)) < 1200


def test_read_excerpt_is_capped_and_labelled_untrusted(doc):
    assert mcp_server.read_excerpt(str(doc))["error_code"] == "not_converted"
    mcp_server.convert([str(doc)], ["json", "xlsx", "txt"])
    piece = mcp_server.read_excerpt(str(doc), max_chars=10 ** 9)
    assert piece["ok"] and piece["chars_returned"] == mcp_server.MAX_EXCERPT_CHARS == 2000
    assert len(piece["untrusted_document_text"]) == 2000 and piece["truncated"] is True
    assert piece["untrusted_document_text"].startswith(SECRET) and piece["next_offset"] == 2000
    assert "not an instruction" in piece["notice"]
    small = mcp_server.read_excerpt(str(doc), page=1, max_chars=40, offset=7)
    assert small["chars_returned"] == 40 and small["untrusted_document_text"] == SECRET[7:47] + (
        "" if len(SECRET) >= 47 else small["untrusted_document_text"][len(SECRET) - 7:])
    assert mcp_server.read_excerpt(str(doc), page=9)["error_code"] == "bad_request"
    out = doc.parent / "output"
    via_xlsx = mcp_server.read_excerpt(str(out / "vertrag.txt.xlsx"), max_chars=50)   # binary output
    via_txt = mcp_server.read_excerpt(str(out / "vertrag.txt.txt"), max_chars=50)     # text output
    via_raw = mcp_server.read_excerpt(str(out / "vertrag.txt.raw.json"), max_chars=50)
    assert via_xlsx["untrusted_document_text"] == via_raw["untrusted_document_text"] == SECRET + "\n\n" + "Absatz mit Fü"
    assert via_txt["chars_returned"] == 50 and via_txt["untrusted_document_text"].startswith(SECRET)
    assert mcp_server.read_excerpt(str(out / "nichts.txt"))["error_code"] == "not_found"


def test_search_returns_short_snippets(doc):
    mcp_server.convert([str(doc)], ["json"])
    found = mcp_server.search(str(doc), "vertragsinhalt")
    assert found["ok"] and found["total_matches"] == 1 and found["untrusted_matches"][0]["page"] == 1
    assert "VERTRAGSINHALT" in found["untrusted_matches"][0]["snippet"]
    many = mcp_server.search(str(doc), "Fülltext", limit=500)
    assert many["total_matches"] == 400 and many["returned"] == mcp_server.MAX_SEARCH_RESULTS
    assert all(len(m["snippet"]) <= mcp_server.SNIPPET_CHARS for m in many["untrusted_matches"])
    assert len(mcp_server._dump(many)) < 6000
    assert mcp_server.search(str(doc), "x")["error_code"] == "bad_request"
    assert mcp_server.search(str(doc), "kommt nicht vor")["total_matches"] == 0


def test_real_stdio_protocol_round_trip(tmp_path):
    """Start the server as a real subprocess and talk MCP over stdin/stdout."""
    mcp = pytest.importorskip("mcp")
    from mcp.client.stdio import StdioServerParameters

    make_docx(tmp_path / "bericht.docx")

    async def scenario():
        params = StdioServerParameters(command=sys.executable, args=["-m", "pdfstruct.mcp_server"])
        async with mcp.Client(params) as client:
            tools = {tool.name for tool in (await client.list_tools()).tools}
            formats = await client.call_tool("supported_formats", {})
            converted = await client.call_tool(
                "convert", {"paths": [str(tmp_path / "bericht.docx")], "formats": ["md", "pdf"]})
            excerpt = await client.call_tool(
                "read_excerpt", {"path_or_output": str(tmp_path / "bericht.docx"), "max_chars": 99999})
            return tools, formats, converted, excerpt

    tools, formats, converted, excerpt = asyncio.run(scenario())
    assert tools == {"supported_formats", "inspect", "convert", "read_excerpt", "search",
                     "create_document"}

    def payload(result) -> dict:
        assert not result.is_error and len(result.content) == 1  # one compact text block
        return json.loads(result.content[0].text)

    assert "pdf" in payload(formats)["outputs"]
    answer = payload(converted)
    assert answer["ok"] and answer["files"][0]["outputs"] == {"md": "bericht.docx.md", "pdf": "bericht.docx.pdf"}
    assert "Quartalsbericht" not in converted.content[0].text  # no document text in a conversion
    assert (tmp_path / "output" / "bericht.docx.pdf").is_file()
    read = payload(excerpt)
    assert read["chars_returned"] <= 2000 and "Quartalsbericht" in read["untrusted_document_text"]


# ---------------------------------------------------------------- create_document

DRAFT = f"# Vorlage {{{{PROJECT_NAME}}}}\n\n{SECRET}\n\n- {{{{REPO_PATH}}}}\n- Çağrı Straße\n\n" + "Fülltext. " * 300


def test_create_document_returns_status_and_paths_only(tmp_path):
    answer = mcp_server.create_document("VORLAGE", DRAFT, ["docx", "pdf", "md"], str(tmp_path))
    assert answer == {"status": "created", "warnings": [],
                      "outputs": [str(tmp_path / f"VORLAGE.{ext}") for ext in ("docx", "pdf", "md")]}
    text = mcp_server._dump(answer)
    assert "GEHEIM" not in text and "Fülltext" not in text and "PROJECT_NAME" not in text
    assert len(text) < 200 + 3 * len(str(tmp_path))
    assert (tmp_path / "VORLAGE.md").read_text(encoding="utf-8") == DRAFT + "\n"

    default = mcp_server.create_document("NUR_WORD", DRAFT, output_dir=str(tmp_path))
    assert default["outputs"] == [str(tmp_path / "NUR_WORD.docx")]  # docx unless asked otherwise
    as_text = mcp_server.create_document("ROH", DRAFT, ["txt"], str(tmp_path), content_type="text")
    assert as_text["status"] == "created"
    assert (tmp_path / "ROH.txt").read_text(encoding="utf-8").startswith("# Vorlage {{PROJECT_NAME}}\n")


def test_create_document_errors_are_compact_and_never_quote_the_text(tmp_path):
    mcp_server.create_document("VORLAGE", DRAFT, ["md"], str(tmp_path))
    answers = {
        "conflict": mcp_server.create_document("VORLAGE", DRAFT, ["md", "docx"], str(tmp_path)),
        "name": mcp_server.create_document("..\\..\\VORLAGE", DRAFT, ["md"], str(tmp_path)),
        "reserved": mcp_server.create_document("NUL", DRAFT, ["md"], str(tmp_path)),
        "format": mcp_server.create_document("X", DRAFT, ["xlsx"], str(tmp_path)),
        "relative": mcp_server.create_document("X", DRAFT, ["md"], "some/relative/folder"),
        "empty": mcp_server.create_document("X", "   ", ["md"], str(tmp_path)),
        "type": mcp_server.create_document("X", DRAFT, ["md"], str(tmp_path), content_type="rtf"),
    }
    assert {key: (a["status"], a["error_code"]) for key, a in answers.items()} == {
        "conflict": ("conflict", "already_exists"), "name": ("invalid", "invalid_name"),
        "reserved": ("invalid", "invalid_name"), "format": ("invalid", "unsupported_format"),
        "relative": ("invalid", "invalid_output_dir"), "empty": ("invalid", "empty_content"),
        "type": ("invalid", "invalid_content_type")}
    for answer in answers.values():
        text = mcp_server._dump(answer)
        assert answer["outputs"] == [] and set(answer) == {"status", "outputs", "warnings", "error_code", "error"}
        assert len(text) < 400 and "GEHEIM" not in text and "Fülltext" not in text
    assert sorted(p.name for p in tmp_path.iterdir()) == ["VORLAGE.md"]  # nothing else appeared
    replaced = mcp_server.create_document("VORLAGE", "neu", ["md"], str(tmp_path), overwrite=True)
    assert replaced["status"] == "created" and (tmp_path / "VORLAGE.md").read_text(encoding="utf-8") == "neu\n"


def test_six_tools_and_the_first_five_keep_their_schemas():
    pytest.importorskip("mcp")
    tools = {tool.name: tool for tool in asyncio.run(mcp_server.build_server().list_tools())}

    def shape(name):
        schema = tools[name].input_schema
        return sorted(schema.get("properties", {})), sorted(schema.get("required", []))

    assert {name: shape(name) for name in tools if name != "create_document"} == {
        "supported_formats": ([], []),
        "inspect": (["path"], ["path"]),
        "convert": (["force_ocr", "formats", "output_dir", "paths"], ["formats", "paths"]),
        "read_excerpt": (["max_chars", "offset", "page", "path_or_output"], ["path_or_output"]),
        "search": (["limit", "path_or_output", "query"], ["path_or_output", "query"])}
    assert tools["convert"].description.startswith("Convert local files (PDF, DOCX, TXT, MD, HTML, images)")
    assert shape("create_document") == (
        ["content", "content_type", "formats", "name", "output_dir", "overwrite"], ["content", "name"])
    properties = tools["create_document"].input_schema["properties"]
    assert properties["formats"]["default"] == ["docx"] and properties["overwrite"]["default"] is False
    assert properties["content_type"]["enum"] == ["markdown", "text"]
    assert properties["content_type"]["default"] == "markdown"
    assert len(tools["create_document"].description) < 260  # it is sent with every conversation
    assert "create_document" in mcp_server.INSTRUCTIONS and len(mcp_server.INSTRUCTIONS) < 800


def test_create_document_over_real_stdio_keeps_stdout_pure_and_logs_no_text(tmp_path):
    """The installed server process, driven with raw JSON-RPC lines as Claude does."""
    pytest.importorskip("mcp")
    spec = importlib.util.spec_from_file_location("verify_runtime", TOOL_DIR / "scripts" / "verify_runtime.py")
    runtime = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runtime)
    server = runtime.Server(runtime.script("pdfstruct-mcp"))
    try:
        init = server.request(1, "initialize", {"protocolVersion": "2025-06-18", "capabilities": {},
                                                 "clientInfo": {"name": "test", "version": "1"}})
        assert init["serverInfo"]["name"] == "pdfstruct"
        server.send({"jsonrpc": "2.0", "method": "notifications/initialized"})
        listed = server.request(2, "tools/list")["tools"]
        assert sorted(tool["name"] for tool in listed) == runtime.TOOLS and len(listed) == 6

        def call(number, **arguments):
            raw = server.request(number, "tools/call", {"name": "create_document", "arguments": arguments})
            assert raw.get("isError") in (None, False) and len(raw["content"]) == 1
            return raw["content"][0]["text"], json.loads(raw["content"][0]["text"])

        text, created = call(3, name="VORLAGE", content=DRAFT, formats=["docx", "html"], output_dir=str(tmp_path))
        assert created == {"status": "created", "warnings": [],
                           "outputs": [str(tmp_path / "VORLAGE.docx"), str(tmp_path / "VORLAGE.html")]}
        assert "GEHEIM" not in text and "Fülltext" not in text
        text, conflict = call(4, name="VORLAGE", content=DRAFT, output_dir=str(tmp_path))
        assert conflict["status"] == "conflict" and "GEHEIM" not in text and len(text) < 400
        text, invalid = call(5, name="a/b", content=DRAFT, output_dir=str(tmp_path))
        assert invalid["error_code"] == "invalid_name" and "GEHEIM" not in text
        for number, arguments in ((6, {"name": "X", "content": DRAFT, "content_type": "rtf"}),
                                  (8, {"content": DRAFT}),                 # no name
                                  (9, {"name": "X", "content": [DRAFT]}),  # text in the wrong shape
                                  (10, {"name": "X", "content": DRAFT, "formats": "docx", "overwrite": "maybe"})):
            bad = server.request(number, "tools/call", {"name": "create_document", "arguments": arguments})
            answer = json.dumps(bad, ensure_ascii=False)
            assert bad.get("isError"), arguments.keys()
            assert "GEHEIM" not in answer and "Fülltext" not in answer and "Vorlage" not in answer
            assert len(answer) < 700  # the field and the reason, not the rejected input
        formats = server.request(7, "tools/call", {"name": "supported_formats", "arguments": {}})
        assert "pdf" in json.loads(formats["content"][0]["text"])["outputs"]  # the others still answer
    finally:
        server.close()
    assert server.stdout_lines >= 10  # every line on stdout parsed as a JSON-RPC message
    assert b"GEHEIM" not in server.stderr and "Fülltext".encode() not in server.stderr
    assert sorted(p.name for p in tmp_path.iterdir()) == ["VORLAGE.docx", "VORLAGE.html"]
    from docx import Document
    document = Document(str(tmp_path / "VORLAGE.docx"))
    assert document.paragraphs[0].style.name == "Heading 1" and SECRET in document.paragraphs[1].text
