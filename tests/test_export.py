import csv
import json
import os
import sqlite3
import subprocess
from pathlib import Path

import pytest

import pdf_export
import pdf_to_json
from conftest import TOOL_DIR, UMLAUT_LINE, make_scan_pdf, make_text_pdf

UMLAUTS = "äöüÄÖÜß"


def export(fmt, src, out, *extra) -> int:
    return pdf_export.main(["--format", fmt, "--input", str(src), "--output", str(out),
                            *map(str, extra)])


def forbid_extraction(monkeypatch):
    def boom(*args, **kwargs):
        raise AssertionError("extraction must not run when raw.json exists")
    monkeypatch.setattr(pdf_to_json, "process_pdf", boom)
    monkeypatch.setattr(pdf_to_json, "main", boom)


@pytest.fixture
def project(tmp_path):
    """A project folder with pdf\\doc.pdf (5 pages, umlauts) already extracted to output."""
    src, out = tmp_path / "pdf", tmp_path / "output"
    src.mkdir()
    make_text_pdf(src / "doc.pdf", pages=5)
    assert pdf_to_json.main(["--input", str(src), "--output", str(out)]) == 0
    return src, out


def sqlite_counts(path: Path) -> tuple:
    conn = sqlite3.connect(str(path))
    try:
        return tuple(conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                     for t in ("documents", "pages", "blocks"))
    finally:
        conn.close()


def test_all_formats_from_existing_raw(project, monkeypatch):
    src, out = project
    raw_file = out / "doc.raw.json"
    raw = json.loads(raw_file.read_text(encoding="utf-8"))
    block_total = sum(len(p["blocks"]) for p in raw["pages"])
    raw_mtime = raw_file.stat().st_mtime_ns
    forbid_extraction(monkeypatch)

    for fmt in pdf_export.FORMATS:
        assert export(fmt, src, out) == 0
    for suffix in pdf_export.FORMATS.values():
        assert (out / f"doc{suffix}").stat().st_size > 0
    assert raw_file.stat().st_mtime_ns == raw_mtime
    assert not list(out.glob("*.tmp"))

    txt = (out / "doc.txt").read_text(encoding="utf-8")
    assert UMLAUT_LINE in txt
    assert txt.index("===== Page 1 / 5 =====") < txt.index("Seite 1 -") \
        < txt.index("===== Page 2 / 5 =====") < txt.index("Seite 5 -")

    md = (out / "doc.md").read_text(encoding="utf-8")
    assert md.startswith("# doc.pdf\n") and md.count("## Page ") == 5
    assert UMLAUT_LINE in md and "|" not in md  # no invented tables
    assert md.index("Zeile 1:") < md.index("Zeile 2:") < md.index("Zeile 8:")

    page = (out / "doc.html").read_text(encoding="utf-8")
    assert '<meta charset="utf-8">' in page and UMLAUT_LINE in page
    assert page.count("<section ") == 5 and 'id="p5"' in page
    assert "method: native" in page and "score: 1.00" in page and "bbox" in page
    assert "http://" not in page and "https://" not in page and "<script" not in page

    with (out / "doc.csv").open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert list(rows[0]) == pdf_export.BLOCK_COLUMNS
    assert len(rows) == block_total
    assert rows[0]["source_file"] == "doc.pdf" and rows[0]["method"] == "native"
    assert rows[0]["page"] == "1" and rows[0]["block_index"] == "0"
    assert float(rows[0]["x1"]) < float(rows[0]["x2"]) and rows[0]["confidence"] == ""
    assert any(r["text"] == UMLAUT_LINE for r in rows)

    records = [json.loads(line) for line in
               (out / "doc.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(records) == block_total
    assert records[0]["page"] == 1 and len(records[0]["bbox"]) == 4
    assert any(r["text"] == UMLAUT_LINE for r in records)
    assert UMLAUTS[0] in (out / "doc.jsonl").read_text(encoding="utf-8")  # not \u-escaped

    from openpyxl import load_workbook
    workbook = load_workbook(out / "doc.xlsx")
    assert workbook.sheetnames == ["Pages", "Blocks"]
    pages, blocks = workbook["Pages"], workbook["Blocks"]
    assert [c.value for c in blocks[1]] == pdf_export.BLOCK_COLUMNS
    assert pages.max_row == 6 and blocks.max_row == block_total + 1
    assert UMLAUT_LINE in pages.cell(row=2, column=pdf_export.PAGE_COLUMNS.index("text") + 1).value
    assert any(row[3] == UMLAUT_LINE for row in blocks.iter_rows(min_row=2, values_only=True))
    assert blocks.column_dimensions["D"].width >= 60 and blocks.freeze_panes == "A2"

    from docx import Document
    document = Document(str(out / "doc.docx"))
    texts = [p.text for p in document.paragraphs]
    assert [t for t in texts if t.startswith("Page ")] == [f"Page {n}" for n in range(1, 6)]
    assert any(UMLAUT_LINE in t for t in texts)
    assert document.element.xml.count('w:type="page"') == 4

    assert sqlite_counts(out / "doc.sqlite") == (1, 5, block_total)
    conn = sqlite3.connect(str(out / "doc.sqlite"))
    try:
        row = conn.execute("SELECT source_file, page, text, x1, y1, x2, y2, method, confidence"
                           " FROM blocks WHERE text = ?", (UMLAUT_LINE,)).fetchone()
        assert row[0] == "doc.pdf" and row[1] == 1 and row[3] < row[5] and row[7] == "native"
        assert UMLAUT_LINE in conn.execute("SELECT text FROM pages WHERE page = 3").fetchone()[0]
    finally:
        conn.close()


def test_second_run_is_safe(project, monkeypatch):
    src, out = project
    forbid_extraction(monkeypatch)
    for fmt in pdf_export.FORMATS:
        assert export(fmt, src, out) == 0
    first = {s: (out / f"doc{s}").read_bytes() for s in (".html", ".txt", ".md", ".csv", ".jsonl")}
    counts = sqlite_counts(out / "doc.sqlite")
    for fmt in pdf_export.FORMATS:
        assert export(fmt, src, out) == 0
    for suffix, content in first.items():
        assert (out / f"doc{suffix}").read_bytes() == content
    assert sqlite_counts(out / "doc.sqlite") == counts  # replaced, not duplicated
    assert not list(out.glob("*.tmp"))


def test_works_without_the_pdf(project, monkeypatch, capsys):
    src, out = project
    (src / "doc.pdf").unlink()
    src.rmdir()
    forbid_extraction(monkeypatch)
    assert export("txt", src, out) == 0
    assert UMLAUT_LINE in (out / "doc.txt").read_text(encoding="utf-8")
    printed = capsys.readouterr().out
    assert "RAW_REUSED: 1" in printed and "EXTRACTED: 0" in printed


def test_missing_raw_runs_pipeline_once(tmp_path, capsys):
    src, out = tmp_path / "pdf", tmp_path / "output"
    src.mkdir()
    make_text_pdf(src / "a.pdf")
    make_text_pdf(src / "b.pdf", pages=2)
    assert export("md", src, out) == 0
    assert (out / "a.raw.json").is_file() and (out / "b.md").is_file()
    assert "EXTRACTED: 2" in capsys.readouterr().out

    # a new PDF arrives: only that one is extracted, the others keep their raw.json
    make_text_pdf(src / "c.pdf")
    mtime = (out / "a.raw.json").stat().st_mtime_ns
    assert export("txt", src, out, "--native-only") == 0
    printed = capsys.readouterr().out
    assert "EXTRACTED: 1" in printed and "RAW_REUSED: 2" in printed
    assert (out / "a.raw.json").stat().st_mtime_ns == mtime
    assert (out / "c.txt").is_file()


def test_ocr_pdf_exports_confidence(tmp_path, ocr_available):
    if not ocr_available:
        pytest.skip("no OCR python available")
    src, out = tmp_path / "pdf", tmp_path / "output"
    src.mkdir()
    make_scan_pdf(src / "scan.pdf")
    for fmt in pdf_export.FORMATS:
        assert export(fmt, src, out) == 0
    assert json.loads((out / "scan.raw.json").read_text(encoding="utf-8"))["ocr_used"] is True
    with (out / "scan.csv").open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows and all(r["method"] == "ocr" and 0 < float(r["confidence"]) <= 1 for r in rows)
    assert "Hallo" in (out / "scan.txt").read_text(encoding="utf-8")
    page = (out / "scan.html").read_text(encoding="utf-8")
    assert "method: ocr" in page and "confidence:" in page
    conn = sqlite3.connect(str(out / "scan.sqlite"))
    try:
        assert conn.execute("SELECT COUNT(*) FROM blocks WHERE method = 'ocr'"
                            " AND confidence IS NOT NULL").fetchone()[0] == len(rows)
    finally:
        conn.close()


def test_control_chars_and_formula_text(tmp_path):
    raw = {"schema_version": "1.0", "source": {"file_name": "x.pdf"}, "source_file": "x.pdf",
           "page_count": 1, "extraction_method": "native", "text_layer_score": 1.0,
           "ocr_used": False, "review_pages": [],
           "pages": [{"page": 1, "width": 100, "height": 100, "rotation": 0, "method": "native",
                      "text_layer": {"score": 1.0}, "ocr_confidence": None, "needs_review": False,
                      "warnings": [],
                      "text": "=SUM(A1)\nBei\x02\nspiel \x0b<b>ß</b>\n# kein Titel",
                      "blocks": [{"text": "=SUM(A1)", "bbox": [1, 2, 3, 4], "order": 0},
                                 {"text": "Bei\x02", "bbox": [1, 5, 3, 6], "order": 1}]}]}
    for fmt, exporter in pdf_export.EXPORTERS.items():
        exporter(raw, tmp_path / f"x{pdf_export.FORMATS[fmt]}")
    from openpyxl import load_workbook
    blocks = load_workbook(tmp_path / "x.xlsx")["Blocks"]
    assert blocks["D2"].value == "=SUM(A1)" and blocks["D2"].data_type == "s"
    assert blocks["D3"].value == "Bei-"
    assert "Bei-\nspiel <b>ß</b>" in (tmp_path / "x.txt").read_text(encoding="utf-8")
    assert "\\# kein Titel" in (tmp_path / "x.md").read_text(encoding="utf-8")
    page = (tmp_path / "x.html").read_text(encoding="utf-8")
    assert "=SUM(A1)" in page and "Bei-" in page and "\x02" not in page

    raw["pages"][0]["blocks"] = []  # no blocks -> page text is shown, HTML-escaped
    pdf_export.export_html(raw, tmp_path / "x.html")
    assert "&lt;b&gt;ß&lt;/b&gt;" in (tmp_path / "x.html").read_text(encoding="utf-8")


def test_nothing_to_do(tmp_path, capsys):
    assert export("txt", tmp_path / "pdf", tmp_path / "output") == 2
    assert "ERROR" in capsys.readouterr().out


@pytest.mark.skipif(os.name != "nt", reason="cmd wrappers are Windows only")
def test_cmd_wrappers(project):
    src, out = project
    for fmt in pdf_export.FORMATS:
        result = subprocess.run(["cmd", "/c", str(TOOL_DIR / f"pdf{fmt}.cmd")],
                                cwd=src.parent, capture_output=True, text=True)
        assert result.returncode == 0, result.stdout + result.stderr
        assert f"FORMAT: {fmt}" in result.stdout and "EXTRACTED: 0" in result.stdout
        assert (out / f"doc{pdf_export.FORMATS[fmt]}").is_file()
