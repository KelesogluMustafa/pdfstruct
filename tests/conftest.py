"""Test fixtures: small PDFs generated on the fly (reportlab + Pillow)."""
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from pdfstruct import extract as pdf_to_json

TOOL_DIR = Path(__file__).resolve().parent.parent  # repo root

UMLAUT_LINE = "Die Straße über den Fluss: ä ö ü Ä Ö Ü ß"


def make_text_pdf(path: Path, pages: int = 1, blank_pages: tuple = ()) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    for number in range(1, pages + 1):
        if number not in blank_pages:
            c.setFont("Helvetica", 12)
            c.drawString(72, 770, f"Seite {number} - Beispieltext")
            c.drawString(72, 750, UMLAUT_LINE)
            for row in range(8):
                c.drawString(72, 720 - row * 18, f"Zeile {row + 1}: Das ist ein normaler Satz mit Text.")
        c.showPage()
    c.save()
    return path


def _scan_image() -> Image.Image:
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default(size=44)
    for row, line in enumerate(["Gescannte Seite", "Hallo Welt Beispiel", "Rechnung Nummer 12345"]):
        draw.text((120, 200 + row * 110), line, fill="black", font=font)
    return image


def make_scan_pdf(path: Path) -> Path:
    _scan_image().save(str(path), "PDF", resolution=150.0)
    return path


def make_mixed_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica", 12)
    c.drawString(72, 770, "Native Seite mit Textebene")
    c.drawString(72, 750, UMLAUT_LINE)
    c.showPage()
    c.drawImage(ImageReader(_scan_image()), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.save()
    return path


@pytest.fixture
def run_tool():
    def _run(*argv) -> int:
        return pdf_to_json.main([str(a) for a in argv])
    return _run


@pytest.fixture(scope="session")
def ocr_available() -> bool:
    from pdfstruct import ocr
    return ocr.available()


# ---------------------------------------------------------------- non-PDF fixtures

UNICODE_LINE = "Çağrı Öğretmen İstanbul ığüşöç – Straße größer Ärger"


def make_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Quartalsbericht", level=1)
    document.add_paragraph(UNICODE_LINE)
    document.add_heading("Aufgaben", level=2)
    document.add_paragraph("Erster Punkt", style="List Bullet")
    document.add_paragraph("Zweiter Punkt", style="List Bullet")
    document.add_paragraph("Schritt eins", style="List Number")
    table = document.add_table(rows=2, cols=3)
    for r, row in enumerate([["Name", "Menge", "Preis"], ["Apfel", "3", "1,20"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("Schlussabsatz nach der Tabelle.")
    document.save(str(path))
    return path


def make_image(path: Path, lines=("Gescannte Seite", "Hallo Welt Beispiel", "Rechnung Nummer 12345"),
               size=(1240, 700), **save_options) -> Path:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default(size=44)
    for row, line in enumerate(lines):
        draw.text((80, 90 + row * 110), line, fill="black", font=font)
    image.save(str(path), **save_options)
    return path


def make_tiff(path: Path, pages=("Erste Seite Alpha", "Zweite Seite Beta")) -> Path:
    frames = []
    for text in pages:
        frame = Image.new("RGB", (1000, 500), "white")
        draw = ImageDraw.Draw(frame)
        try:
            font = ImageFont.truetype("arial.ttf", 48)
        except OSError:
            font = ImageFont.load_default(size=48)
        draw.text((80, 180), text, fill="black", font=font)
        frames.append(frame)
    frames[0].save(str(path), save_all=True, append_images=frames[1:])
    return path
