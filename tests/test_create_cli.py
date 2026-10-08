"""pdfstruct-create: text from a file or standard input, through the shared create service."""
import io
import subprocess
import sys

from docx import Document

from pdfstruct import cli
from test_create import MARKDOWN, PLACEHOLDERS, TEXT, UNICODE


def test_cli_reads_a_content_file(tmp_path, capsys):
    source = tmp_path / "şablon notları.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    assert cli.create(["--content-file", source.name, "--name", "TEMPLATE", "--format", "docx,pdf",
                       "--format", "md", "--output", "out"], cwd=tmp_path) == 0
    out = capsys.readouterr().out
    assert out.splitlines() == [f"CREATED: {tmp_path / 'out' / ('TEMPLATE.' + ext)}" for ext in ("docx", "pdf", "md")]
    assert "PROJECT_NAME" not in out and "Intro" not in out  # paths only
    assert Document(str(tmp_path / "out" / "TEMPLATE.docx")).paragraphs[0].style.name == "Heading 1"
    assert (tmp_path / "out" / "TEMPLATE.md").read_text(encoding="utf-8") == MARKDOWN

    assert cli.create(["--content-file", source.name, "--name", "TEMPLATE", "--output", "out"], cwd=tmp_path) == 1
    assert capsys.readouterr().out.startswith("CONFLICT: already exists, nothing was written: TEMPLATE.docx")
    assert cli.create(["--content-file", source.name, "--name", "TEMPLATE", "--output", "out",
                       "--overwrite"], cwd=tmp_path) == 0
    capsys.readouterr()

    assert cli.create(["--content-file", source.name], cwd=tmp_path) == 0  # name and folder from the file
    assert (tmp_path / "şablon notları.docx").is_file()
    capsys.readouterr()
    assert cli.create(["--content-file", source.name, "--format", "md", "--overwrite"], cwd=tmp_path) == 2
    assert "is the text file itself" in capsys.readouterr().out and source.read_text(encoding="utf-8") == MARKDOWN

    note = tmp_path / "note.txt"
    note.write_text(TEXT, encoding="utf-8-sig")
    assert cli.create(["--content-file", "note.txt", "--format", "html"], cwd=tmp_path) == 0  # .txt: plain text
    assert "<h1>" not in (tmp_path / "note.html").read_text(encoding="utf-8")
    assert cli.create(["--content-file", "note.txt", "--format", "html", "--type", "markdown",
                       "--name", "as_md"], cwd=tmp_path) == 0
    assert "<h1>not a heading</h1>" in (tmp_path / "as_md.html").read_text(encoding="utf-8")
    capsys.readouterr()

    for argv, expected in ((["--content-file", "missing.md"], "ERROR: cannot read"),
                           (["--content-file", "note.txt", "--name", "..\\x"], "INVALID: name must be"),
                           (["--content-file", "note.txt", "--format", "xlsx"], "INVALID: unsupported format: xlsx")):
        assert cli.create(argv, cwd=tmp_path) == 2
        assert expected in capsys.readouterr().out


def test_cli_reads_standard_input_as_utf8(tmp_path):
    def run(*argv, data):
        return subprocess.run([sys.executable, "-c", "import sys; from pdfstruct.cli import create; sys.exit(create())",
                               *argv], input=data, cwd=tmp_path, capture_output=True, timeout=120)
    done = run("--name", "FROM_STDIN", "--format", "txt,docx,html", data=MARKDOWN.encode("utf-8"))
    out = done.stdout.decode("utf-8")
    assert done.returncode == 0 and done.stderr == b"", done.stderr
    assert out.splitlines() == [f"CREATED: {tmp_path / ('FROM_STDIN.' + ext)}" for ext in ("txt", "docx", "html")]
    written = (tmp_path / "FROM_STDIN.txt").read_text(encoding="utf-8")
    assert UNICODE in written and all(p in written for p in PLACEHOLDERS) and UNICODE not in out
    assert UNICODE in "\n".join(p.text for p in Document(str(tmp_path / "FROM_STDIN.docx")).paragraphs)

    assert run("--format", "txt", data=b"text").returncode == 2             # no --name
    empty = run("--name", "EMPTY", data=b"  \n")
    assert empty.returncode == 2 and b"INVALID: content is empty" in empty.stdout
    again = run("--name", "FROM_STDIN", "--format", "txt", data=b"other")
    assert again.returncode == 1 and (tmp_path / "FROM_STDIN.txt").read_text(encoding="utf-8") == written
    help_text = run("--help", data=b"").stdout.decode("utf-8")
    assert "--content-file" in help_text and "standard input" in help_text


def test_cli_refuses_an_oversized_file_without_loading_it(tmp_path, capsys, monkeypatch):
    from pdfstruct import create
    monkeypatch.setattr(create, "MAX_CONTENT_CHARS", 10)
    (tmp_path / "big.md").write_bytes(b"x" * 1000)
    assert cli.create(["--content-file", "big.md", "--name", "BIG"], cwd=tmp_path) == 2
    assert "larger than 10 characters" in capsys.readouterr().out
    assert sorted(p.name for p in tmp_path.iterdir()) == ["big.md"]


def test_cli_without_text_explains_itself(tmp_path, capsys):
    class Terminal(io.StringIO):
        def isatty(self):
            return True
    assert cli.create(["--name", "X"], cwd=tmp_path, stdin=Terminal()) == 2
    assert "Use --content-file FILE or pipe the text in" in capsys.readouterr().out
    assert list(tmp_path.iterdir()) == []
